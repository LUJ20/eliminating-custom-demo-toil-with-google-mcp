"""Interactive Google Slides / PPTX Presentation Player for Streamlit UI.
Parses native PowerPoint shapes, text, tables, pictures, connectors and presenter notes directly from the .pptx
file and renders an interactive presentation player with slide navigation, notes, and
direct links to Google Slides.
"""
import base64
import io
import os
import json
import html as _html
from typing import Optional

try:
    import pptx
    from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
    # arrow auto-shapes (the deck's architecture flow) -> the glyph the player draws for them
    ARROW_GLYPHS = {MSO_SHAPE.RIGHT_ARROW: "➜", MSO_SHAPE.LEFT_ARROW: "⬅", MSO_SHAPE.DOWN_ARROW: "⬇",
                    MSO_SHAPE.UP_ARROW: "⬆"}
    LINE_TYPE = MSO_SHAPE_TYPE.LINE  # connectors (the stage -> stage and stage -> output lines) are SVG lines
    PICTURE_TYPE = MSO_SHAPE_TYPE.PICTURE  # the template's cover art, logo and panels, the diagram's product icons
except ImportError:
    pptx = None
    ARROW_GLYPHS = {}
    LINE_TYPE = None
    PICTURE_TYPE = None

MAX_PICTURE_PX = 1000  # pictures are inlined as data URIs; anything larger is downscaled first (Pillow, when present)


def picture_data_uri(blob: bytes, content_type: str) -> str:
    """A data: URI for a picture blob, downscaled to MAX_PICTURE_PX on its long side when Pillow can do it (the
    template's cover art is 1200 px and 155 KB; the icons are small already). Without Pillow the bytes go as is."""
    try:
        from PIL import Image
        with Image.open(io.BytesIO(blob)) as im:
            if max(im.size) > MAX_PICTURE_PX:
                im.thumbnail((MAX_PICTURE_PX, MAX_PICTURE_PX))
                out = io.BytesIO()
                im.save(out, format="PNG", optimize=True)
                blob, content_type = out.getvalue(), "image/png"
    except Exception:  # no Pillow, or a format it cannot read: the original bytes
        pass
    return f"data:{content_type or 'image/png'};base64,{base64.b64encode(blob).decode('ascii')}"


def render_presentation_player(deck_path: str, gslides_url: Optional[str] = None, height: int = 520) -> str:
    """Builds an interactive Google Slides HTML player string from a .pptx file."""
    if not os.path.exists(deck_path) or pptx is None:
        return "<p style='color:#5F6368;'>Presentation deck preview not available.</p>"

    try:
        prs = pptx.Presentation(deck_path)
    except Exception as e:
        return f"<p style='color:#EA4335;'>Failed to load presentation: {_html.escape(str(e))}</p>"

    sw = float(prs.slide_width or 12192000)
    sh = float(prs.slide_height or 6858000)
    slides_html = []
    notes_list = []
    total_slides = len(prs.slides)

    for s_idx, slide in enumerate(prs.slides):
        shape_divs = []
        svg_lines = []  # connectors, drawn once per slide in an overlay
        slide_title = f"Slide {s_idx + 1}"
        first_title_found = False

        for order, shp in enumerate(slide.shapes, 1):  # document order is the stacking order, as PowerPoint paints
            z = order + 1
            if LINE_TYPE is not None and getattr(shp, "shape_type", None) == LINE_TYPE:
                try:
                    x1, y1 = shp.begin_x / sw * 100.0, shp.begin_y / sh * 100.0
                    x2, y2 = shp.end_x / sw * 100.0, shp.end_y / sh * 100.0
                    color = f"#{shp.line.color.rgb}" if shp.line.color and shp.line.color.rgb else "#9AA0A6"
                    dash = ' stroke-dasharray="6 4"' if shp.line.dash_style else ""
                    attrs = f'stroke="{color}" stroke-width="1.5" fill="none" vector-effect="non-scaling-stroke"{dash}'
                    if "bentConnector" in (shp._element.xpath("string(.//a:prstGeom/@prst)") or ""):
                        my = (y1 + y2) / 2  # an elbow: down to the midpoint, across, down
                        svg_lines.append(f'<path d="M{x1:.2f} {y1:.2f} V{my:.2f} H{x2:.2f} V{y2:.2f}" {attrs}/>')
                    else:
                        svg_lines.append(f'<line x1="{x1:.2f}" y1="{y1:.2f}" x2="{x2:.2f}" y2="{y2:.2f}" {attrs}/>')
                except Exception:
                    pass
                continue
            l_pct = round(max(0.0, float(getattr(shp, "left", 0) or 0) / sw * 100.0), 2)
            t_pct = round(max(0.0, float(getattr(shp, "top", 0) or 0) / sh * 100.0), 2)
            w_pct = round(max(0.5, float(getattr(shp, "width", 0) or 0) / sw * 100.0), 2)
            h_pct = round(max(0.5, float(getattr(shp, "height", 0) or 0) / sh * 100.0), 2)

            # Pictures: inlined, kept in proportion inside their box
            if PICTURE_TYPE is not None and getattr(shp, "shape_type", None) == PICTURE_TYPE:
                try:
                    uri = picture_data_uri(shp.image.blob, shp.image.content_type)
                    shape_divs.append(
                        f'<img src="{uri}" alt="" style="position:absolute;left:{l_pct}%;top:{t_pct}%;width:{w_pct}%;'
                        f'height:{h_pct}%;object-fit:contain;z-index:{z};">'
                    )
                except Exception:
                    pass
                continue

            # Table shapes
            if getattr(shp, "has_table", False):
                try:
                    tbl = shp.table
                    tr_list = []
                    for r_i, row in enumerate(tbl.rows):
                        td_list = []
                        for c_i, cell in enumerate(row.cells):
                            c_txt = (cell.text or "").strip()
                            c_bg = "#1A73E8" if r_i == 0 else ("#F8F9FA" if r_i % 2 == 1 else "#FFFFFF")
                            c_fg = "#FFFFFF" if r_i == 0 else "#202124"
                            c_fw = "700" if (r_i == 0 or c_i == 0) else "400"
                            try:
                                if cell.fill and cell.fill.type == 1 and cell.fill.fore_color and cell.fill.fore_color.rgb:
                                    c_bg = f"#{cell.fill.fore_color.rgb}"
                                    if c_bg.upper() in ("#1A73E8", "#174EA6", "#202124", "#34A853", "#EA4335"):
                                        c_fg = "#FFFFFF"
                            except Exception:
                                pass
                            td_list.append(
                                f'<td style="background:{c_bg};color:{c_fg};font-weight:{c_fw};'
                                f'padding:4px 8px;border:1px solid #DADCE0;font-size:0.95cqw;line-height:1.25;vertical-align:top;">'
                                f'{_html.escape(c_txt)}</td>'
                            )
                        tr_list.append(f"<tr>{''.join(td_list)}</tr>")
                    shape_divs.append(
                        f'<div class="gs-shape" style="position:absolute;left:{l_pct}%;top:{t_pct}%;width:{w_pct}%;height:{h_pct}%;overflow:hidden;z-index:{z};">'
                        f'<table style="width:100%;height:100%;border-collapse:collapse;table-layout:fixed;">{"".join(tr_list)}</table>'
                        f'</div>'
                    )
                    continue
                except Exception:
                    pass

            bg_hex = "transparent"
            try:
                if shp.fill and shp.fill.type == 1 and shp.fill.fore_color and shp.fill.fore_color.rgb:
                    bg_hex = f"#{shp.fill.fore_color.rgb}"
            except Exception:
                pass

            border_css = "none"
            try:
                if shp.line and shp.line.color and shp.line.color.rgb:
                    style = "dashed" if shp.line.dash_style else "solid"  # the demo-output group is dashed
                    border_css = f"1px {style} #{shp.line.color.rgb}"
            except Exception:
                pass

            radius_css = "5px"
            try:
                kind = getattr(shp, "auto_shape_type", None)
                if kind == MSO_SHAPE.OVAL:
                    radius_css = "50%"
                elif kind in ARROW_GLYPHS:
                    # Arrows (the architecture flow) render as a symbol in the arrow's colour, not as a filled box
                    color = bg_hex if bg_hex != "transparent" else "#1A73E8"
                    shape_divs.append(
                        f'<div style="position:absolute;left:{l_pct}%;top:{t_pct}%;width:{w_pct}%;height:{h_pct}%;'
                        f'display:flex;align-items:center;justify-content:center;color:{color};font-size:1.6cqw;'
                        f'font-weight:bold;z-index:{z};">{ARROW_GLYPHS[kind]}</div>'
                    )
                    continue
            except Exception:
                pass

            paras_html = []
            if getattr(shp, "has_text_frame", False):
                for p_obj in shp.text_frame.paragraphs:
                    raw_t = (p_obj.text or "").strip()
                    if not raw_t:
                        continue
                    if not first_title_found and len(raw_t) > 3:
                        slide_title = raw_t[:45]
                        first_title_found = True

                    pt_val = 10.0
                    is_bold = False
                    col_hex = "#202124"
                    fonts = [f for f in (p_obj.font, *(r.font for r in p_obj.runs[:1])) if f is not None]
                    sized = False
                    for f in fonts:
                        try:
                            if f.size:
                                pt_val, sized = float(f.size.pt), True
                                break
                        except Exception:
                            continue
                    if not sized and getattr(shp, "is_placeholder", False):
                        pt_val = 24.0  # a title placeholder inherits its size from the layout (the template's 24 pt)
                    for f in fonts:
                        try:
                            if f.bold:
                                is_bold = True
                                break
                        except Exception:
                            continue
                    for f in fonts:
                        try:
                            if f.color and f.color.type is not None and f.color.rgb:  # a theme colour raises: next
                                col_hex = f"#{f.color.rgb}"
                                break
                        except Exception:
                            continue

                    # 1 cqw = 1% of the slide width; a point is 1/72 in of a 13.333 in slide = 0.104 cqw
                    cqw_size = round(max(0.7, min(8.4, pt_val / 9.6)), 2)
                    spans = []
                    for r in p_obj.runs:
                        if not r.text:
                            continue
                        try:
                            r_bold = r.font.bold
                        except Exception:
                            r_bold = None
                        fw = "700" if (is_bold if r_bold is None else r_bold) else "400"
                        spans.append(f'<span style="font-weight:{fw};">{_html.escape(r.text)}</span>')
                    body = "".join(spans) or _html.escape(raw_t)
                    paras_html.append(
                        f'<div style="font-size:{cqw_size}cqw;font-weight:{"700" if is_bold else "400"};color:{col_hex};'
                        f'line-height:1.22;margin-bottom:0.2cqw;word-break:break-word;">{body}</div>'
                    )

            pad_css, anchor_css = "0", ""
            if paras_html:  # the text frame's own insets and vertical anchor (the template's headers clear their marks)
                try:
                    tf = shp.text_frame
                    mt, mr, mb, ml = (float(v or 0) / sw * 100.0 for v in (tf.margin_top, tf.margin_right,
                                                                           tf.margin_bottom, tf.margin_left))
                    pad_css = f"{mt:.2f}cqw {mr:.2f}cqw {mb:.2f}cqw {ml:.2f}cqw"
                    anchor = str(tf.vertical_anchor or "")
                    if "MIDDLE" in anchor or "BOTTOM" in anchor:
                        anchor_css = ("display:flex;flex-direction:column;justify-content:"
                                      + ("center" if "MIDDLE" in anchor else "flex-end") + ";")
                except Exception:
                    pad_css = "0.5cqw 0.75cqw"
            overflow = "visible" if (bg_hex == "transparent" and border_css == "none") else "hidden"
            shape_divs.append(
                f'<div class="gs-shape" style="position:absolute;left:{l_pct}%;top:{t_pct}%;width:{w_pct}%;height:{h_pct}%;'
                f'background:{bg_hex};border:{border_css};border-radius:{radius_css};padding:{pad_css};{anchor_css}'
                f'box-sizing:border-box;overflow:{overflow};z-index:{z};">'
                f'{"".join(paras_html)}</div>'
            )

        notes_txt = ""
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_txt = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            pass
        notes_list.append(notes_txt if notes_txt else f"Slide {s_idx + 1}: {slide_title}")

        disp_style = "block" if s_idx == 0 else "none"
        slide_inner_html = "".join(shape_divs)
        if svg_lines:  # the connectors, over the shapes and under nothing clickable
            slide_inner_html += ('<svg viewBox="0 0 100 100" preserveAspectRatio="none" style="position:absolute;'
                                 'left:0;top:0;width:100%;height:100%;z-index:9999;pointer-events:none;">'
                                 + "".join(svg_lines) + '</svg>')
        slides_html.append(
            f'<div class="gslide-frame" id="gslide_{s_idx}" style="display:{disp_style};container-type:inline-size;position:relative;width:97%;height:410px;background:#FFFFFF;overflow:hidden;border-radius:4px;box-shadow:0 3px 12px rgba(0,0,0,0.25);">'
            f'{slide_inner_html}'
            f'</div>'
        )

    options_html = "".join(f'<option value="{i}">Slide {i+1} of {total_slides}</option>' for i in range(total_slides))
    # Notes come from model-written text: keep them inert inside the <script> block.
    notes_json = json.dumps(notes_list).replace("<", "\\u003c").replace(">", "\\u003e").replace("&", "\\u0026")
    slides_link = ""
    if (gslides_url or "").startswith("https://docs.google.com/"):  # only a deck that was really published
        slides_link = (f'<a href="{_html.escape(gslides_url, quote=True)}" target="_blank" rel="noopener noreferrer" '
                       'style="text-decoration:none;color:#1A73E8;font-size:13px;font-weight:600;display:flex;'
                       'align-items:center;gap:6px;background:#FFF;border:1px solid #DADCE0;padding:5px 14px;'
                       'border-radius:6px;"><span style="display:inline-block;width:12px;height:12px;'
                       'background:#FBBC04;border-radius:2px;"></span><span>Open in Google Slides &#8599;</span></a>')

    player_html = f"""
    <div style="font-family:'Google Sans',Roboto,sans-serif;width:100%;border:1px solid #DADCE0;border-radius:8px;overflow:hidden;background:#1F1F1F;box-shadow:0 2px 8px rgba(0,0,0,0.12);">
      <div style="position:relative;background:#121212;height:435px;display:flex;align-items:center;justify-content:center;padding:10px 8px;">
        <div style="position:absolute;top:12px;right:16px;z-index:20;">
          <button onclick="toggleNotes()" style="background:rgba(32,33,36,0.9);border:1px solid #747775;color:#FFF;padding:5px 12px;border-radius:6px;cursor:pointer;font-size:12px;display:flex;align-items:center;gap:6px;">
            <span>Speaker Notes</span>
          </button>
        </div>
        <div id="gs_container" style="display:flex;align-items:center;justify-content:center;width:100%;height:100%;">
          {"".join(slides_html)}
        </div>
        <div id="gs_notes" style="display:none;position:absolute;bottom:12px;left:18px;right:18px;background:rgba(32,33,36,0.96);color:#E8EAED;padding:12px 16px;border-radius:6px;font-size:12px;line-height:1.45;z-index:25;border-left:4px solid #FBBC04;max-height:140px;overflow-y:auto;white-space:pre-line;">
          {_html.escape(notes_list[0] if notes_list else "")}
        </div>
      </div>
      <div style="background:#F1F3F4;height:48px;padding:0 16px;display:flex;align-items:center;justify-content:space-between;border-top:1px solid #DADCE0;">
        <div style="display:flex;align-items:center;gap:10px;">
          <button onclick="stepSlide(-1)" style="background:#FFF;border:1px solid #DADCE0;padding:4px 12px;border-radius:999px;cursor:pointer;font-weight:bold;font-size:13px;" title="Previous slide">&#10094;</button>
          <select id="gs_select" onchange="showSlide(parseInt(this.value))" style="background:#FFF;border:1px solid #DADCE0;padding:4px 12px;border-radius:999px;cursor:pointer;font-size:13px;font-weight:600;color:#202124;">
            {options_html}
          </select>
          <button onclick="stepSlide(1)" style="background:#FFF;border:1px solid #DADCE0;padding:4px 12px;border-radius:999px;cursor:pointer;font-weight:bold;font-size:13px;" title="Next slide">&#10095;</button>
        </div>
        <div style="display:flex;align-items:center;gap:8px;">
          {slides_link}
        </div>
      </div>
    </div>
    <script>
      var curSlide = 0;
      var totalSlides = {total_slides};
      var slideNotes = {notes_json};
      var notesVisible = false;

      function showSlide(idx) {{
        if (idx < 0) idx = totalSlides - 1;
        if (idx >= totalSlides) idx = 0;
        curSlide = idx;
        for (var i = 0; i < totalSlides; i++) {{
          var el = document.getElementById('gslide_' + i);
          if (el) el.style.display = (i === curSlide) ? 'block' : 'none';
        }}
        var sel = document.getElementById('gs_select');
        if (sel) sel.value = String(curSlide);
        var nb = document.getElementById('gs_notes');
        if (nb) nb.innerText = slideNotes[curSlide] || '';
      }}

      function stepSlide(delta) {{
        showSlide(curSlide + delta);
      }}

      function toggleNotes() {{
        notesVisible = !notesVisible;
        var nb = document.getElementById('gs_notes');
        if (nb) nb.style.display = notesVisible ? 'block' : 'none';
      }}
    </script>
    """
    return player_html
