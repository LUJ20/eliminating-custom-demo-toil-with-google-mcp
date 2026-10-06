"""The reference architecture deck: the four slides of the Google Cloud reference architecture template (the SS2
bill-of-materials deck), filled in place from the build so the typography, logo, cover art and table styling are the
template's own.

  1. Cover: the customer, the repeatable pattern the BOM writer named, the month ("REFERENCE ARCHITECTURE" kicker)
  2. Reference architecture: the objective and three design goals on the left (40%), the studio's diagram on the
     right (60%): the stages in execution order as editable shapes coloured by capability tier and joined by arrows,
     the demo outputs in a dashed group, a legend
  3. Design considerations: the template's Well-Architected table (Reliability, Cost Optimization, Security: design
     decision and target metric impact), with the Well-Architected review's verdict above it
  4. Applicability criteria: when to use and when to avoid this architecture, three each

The template file is not in the repository (engine/bom_template.py says where it lives and how to fetch it). Without
it the same four slides are drawn on a blank 16:9 deck at the template's positions and fonts, so a build never fails
for want of the file; the deck's core properties say which was used. Every word on a slide or in the speaker notes
comes from the build (the design, the BOM narrative, the review, the scorecard); nothing is invented.
"""
import datetime as _dt
import os
import re
from typing import Dict, List, NamedTuple, Optional, Sequence, Tuple

import pptx
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from engine import manifest
from engine.well_architected import MIN_PILLAR_SCORE as MIN_PILLAR

FONT, FONT_TEXT, FONT_MEDIUM, FONT_TABLE = "Google Sans", "Google Sans Text", "Google Sans Medium", "DM Sans"
BLUE = RGBColor(26, 115, 232)
DARK_BLUE = RGBColor(23, 78, 166)
RED = RGBColor(234, 67, 53)
YELLOW = RGBColor(242, 153, 0)
GREEN = RGBColor(24, 128, 56)
GOOGLE_GREEN = RGBColor(52, 168, 83)
HEADER = RGBColor(32, 33, 36)
BODY = RGBColor(60, 64, 67)
MUTED = RGBColor(95, 99, 104)
LIGHT_BLUE = RGBColor(232, 240, 254)
BORDER = RGBColor(218, 220, 224)
WHITE = RGBColor(255, 255, 255)
FAIL_RED = RGBColor(197, 34, 31)
GREY_LINE = RGBColor(154, 160, 166)
GREY_FILL = RGBColor(241, 243, 244)
LIGHT_GREEN = RGBColor(230, 244, 234)

# capability tier -> (fill, border). Non-AI stages use the default Google blue outline.
TIER_COLORS: Dict[str, Tuple[RGBColor, RGBColor]] = {
    "reasoning": (LIGHT_GREEN, GREEN), "fast": (LIGHT_GREEN, GREEN), "lite": (LIGHT_GREEN, GREEN),
    "live": (LIGHT_BLUE, BLUE),
    "image": (RGBColor(243, 232, 253), RGBColor(147, 52, 230)),
    "image_fast": (RGBColor(243, 232, 253), RGBColor(147, 52, 230)),
    "video": (RGBColor(252, 232, 230), FAIL_RED), "video_fast": (RGBColor(252, 232, 230), FAIL_RED),
    "speech": (RGBColor(254, 247, 224), RGBColor(227, 116, 0)), "music": (RGBColor(254, 247, 224), RGBColor(227, 116, 0)),
    "embedding": (RGBColor(224, 247, 250), RGBColor(0, 131, 143)),
}
DEFAULT_COLORS = (WHITE, BLUE)
TIER_LABELS = {"reasoning": "Gemini reasoning", "fast": "Gemini fast", "lite": "Gemini lite", "live": "Gemini Live",
               "image": "Image", "image_fast": "Image (fast)", "video": "Video", "video_fast": "Video (fast)",
               "speech": "Speech", "music": "Music", "embedding": "Embeddings"}

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
BLANK_LAYOUT = 6
DECK_VERSION = "8"  # the slide layout; stamped into every deck. Bump it when the slides change: saved decks made
                    # by an older layout are then regenerated from the stored result (build_editor.refresh_deck).
                    # 6: product names as the docs use them today (Agent Platform, formerly Vertex AI)
                    # 7: Well-Architected review slide after the scorecard
                    # 8: the four slides of the Google Cloud reference architecture template (the BOM deck)
TEMPLATE_TAG, BLANK_TAG = "google-cloud-reference-architecture-template", "blank-fallback"
MAX_OUTPUT_CARDS = 4
MAX_CELL_CHARS = 240

# The template: slide indices and shape IDs (engine/bom_template.py documents the file; a dump of its shapes is in
# the module docstring there). Slides 0 and 5-8 are the "make a copy" page and the examples: dropped.
T_COVER, T_ARCH, T_CONS, T_APP = 1, 2, 3, 4
T_DROP = (0, 5, 6, 7, 8)
ID_COVER_TITLE, ID_COVER_KICKER = 539, 543
ID_ARCH_TITLE, ID_ARCH_SUMMARY, ID_ARCH_GOALS, ID_ARCH_DIAGRAM = 549, 550, (551, 552, 553), 554
ID_CONS_TITLE, ID_CONS_SUB, ID_CONS_TABLE = 560, 561, 562
ID_APP_TITLE, ID_USE_HEAD, ID_AVOID_HEAD = 568, 571, 572
ID_USES, ID_AVOIDS = (574, 575, 576), (578, 579, 580)
ID_USE_ICONS, ID_AVOID_ICONS, ID_AVOID_MARK = (581, 582, 583), (584, 585, 586), 577
# Geometry (inches) of the template's shapes, used as they are and for the blank fallback.
G_COVER_TITLE = (0.67, 0.93, 8.45, 4.95)
G_COVER_KICKER = (0.67, 4.97, 5.45, 0.34)
G_TITLE = (0.68, 0.25, 12.34, 0.77)
G_SUMMARY = (0.76, 1.19, 4.88, 1.30)
G_GOAL_TOPS, G_GOAL = (2.72, 3.62, 4.52), (0.76, 4.72, 0.82)       # tops; (left, width, height)
G_DIAGRAM = (5.56, 1.39, 7.22, 4.86)
G_CONS_SUB = (0.80, 1.39, 11.81, 0.24)
G_CONS_TABLE = (0.80, 1.63, 11.81, 3.28)
G_HEAD_USE, G_HEAD_AVOID = (1.59, 2.03, 4.85, 0.24), (7.43, 2.03, 4.85, 0.24)
G_ROW_TOPS, G_ROW_H, G_ROW_W = (2.92, 3.52, 4.12), 0.52, 4.60
G_USE_LEFT, G_AVOID_LEFT, G_ICON_DX = 1.59, 7.43, 0.29


class Line(NamedTuple):
    text: str
    size: float = 11
    color: RGBColor = BODY
    bold: bool = False
    link: str = ""
    font: str = FONT


class Cell(NamedTuple):
    text: str
    link: str = ""
    color: Optional[RGBColor] = None
    bold: bool = False


def tier_colors(tier: str) -> Tuple[RGBColor, RGBColor]:
    return TIER_COLORS.get(tier or "", DEFAULT_COLORS)


def _cut(text, limit: int = MAX_CELL_CHARS) -> str:
    text = " ".join(str(text or "").split())
    return text if len(text) <= limit else text[:limit - 1].rstrip() + "…"


def _fill(tf, lines: Sequence[Line], align=PP_ALIGN.LEFT) -> None:
    tf.clear()
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line.text
        run.font.name, run.font.size, run.font.bold = line.font, Pt(line.size), line.bold
        run.font.color.rgb = line.color
        if line.link.startswith("https://"):
            run.hyperlink.address = line.link


def _textbox(slide, left, top, width, height, lines: Sequence[Line]):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    _fill(tf, lines)
    return box


def _card(slide, left, top, width, height, lines: Sequence[Line], fill: RGBColor = WHITE,
          border: RGBColor = BORDER, radius: float = 0.05, margin=Inches(0.1)) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.adjustments[0] = radius
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border
    shape.line.width = Pt(1.2)
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = margin
    tf.margin_top = tf.margin_bottom = min(margin, Inches(0.08))
    _fill(tf, lines)


def _shape(slide, kind, left, top, width, height, color: RGBColor) -> None:
    shape = slide.shapes.add_shape(kind, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False


def _dashed_box(slide, left, top, width, height, color: RGBColor = BORDER) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.adjustments[0] = 0.06
    shape.fill.background()
    shape.line.color.rgb = color
    shape.line.width = Pt(1.2)
    shape.line.dash_style = MSO_LINE.DASH
    shape.shadow.inherit = False


def _connector(slide, x1, y1, x2, y2, color: RGBColor = GREY_LINE, dashed: bool = True) -> None:
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(1.25 if dashed else 1.75)
    if dashed:
        c.line.dash_style = MSO_LINE.DASH
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))


def _set(shape, lines: Sequence[Line], align=PP_ALIGN.LEFT) -> None:
    """Replace the text of a template shape (its box, position and background stay)."""
    _fill(shape.text_frame, lines, align)


def _place(shape, left: float, top: float, width: float, height: float) -> None:
    shape.left, shape.top, shape.width, shape.height = Inches(left), Inches(top), Inches(width), Inches(height)


def _by_id(slide, shape_id: int):
    return next((s for s in slide.shapes if s.shape_id == shape_id), None)


def _remove(shape) -> None:
    el = shape._element
    el.getparent().remove(el)


# ---------------------------------------------------------------------------------------------- the skeleton
class Skeleton:
    """The shapes the four slides are filled into: found in the template by shape ID, or created on blank slides at
    the template's positions when there is no template."""

    def __init__(self, prs, from_template: bool):
        self.prs, self.from_template = prs, from_template
        self.cover = self.arch = self.cons = self.app = None
        self.cover_title = self.arch_title = self.arch_summary = self.cons_title = self.cons_sub = None
        self.cons_table = self.app_title = self.use_head = self.avoid_head = None
        self.goals: list = []
        self.uses: list = []
        self.avoids: list = []
        self.use_icons: list = []
        self.avoid_icons: list = []
        self.diagram = G_DIAGRAM

    @classmethod
    def from_template(cls, path: str) -> "Skeleton":
        prs = pptx.Presentation(path)
        if len(prs.slides) <= max(T_DROP + (T_APP,)):
            raise ValueError(f"the template has {len(prs.slides)} slides; the reference architecture template has 9")
        sk = cls(prs, True)
        sk.cover, sk.arch, sk.cons, sk.app = (prs.slides[i] for i in (T_COVER, T_ARCH, T_CONS, T_APP))
        sk.cover_title = _by_id(sk.cover, ID_COVER_TITLE)
        sk.arch_title, sk.arch_summary = _by_id(sk.arch, ID_ARCH_TITLE), _by_id(sk.arch, ID_ARCH_SUMMARY)
        sk.goals = [_by_id(sk.arch, i) for i in ID_ARCH_GOALS]
        diagram = _by_id(sk.arch, ID_ARCH_DIAGRAM)
        if diagram is not None:  # the "[REFERENCE DIAGRAM]" placeholder: its rectangle is where the diagram goes
            sk.diagram = tuple(Emu(v).inches for v in (diagram.left, diagram.top, diagram.width, diagram.height))
            _remove(diagram)
        sk.cons_title, sk.cons_sub = _by_id(sk.cons, ID_CONS_TITLE), _by_id(sk.cons, ID_CONS_SUB)
        table = _by_id(sk.cons, ID_CONS_TABLE)
        sk.cons_table = table.table if table is not None and table.has_table else None
        sk.app_title = _by_id(sk.app, ID_APP_TITLE)
        sk.use_head, sk.avoid_head = _by_id(sk.app, ID_USE_HEAD), _by_id(sk.app, ID_AVOID_HEAD)
        sk.uses, sk.avoids = [_by_id(sk.app, i) for i in ID_USES], [_by_id(sk.app, i) for i in ID_AVOIDS]
        sk.use_icons, sk.avoid_icons = [_by_id(sk.app, i) for i in ID_USE_ICONS], [_by_id(sk.app, i) for i in ID_AVOID_ICONS]
        missing = [n for n, s in (("cover title", sk.cover_title), ("architecture title", sk.arch_title),
                                  ("summary", sk.arch_summary), ("considerations table", sk.cons_table),
                                  ("applicability title", sk.app_title)) if s is None]
        if missing or any(s is None for s in sk.goals + sk.uses + sk.avoids):
            raise ValueError("the template is missing shapes: " + ", ".join(missing or ["criteria rows"]))
        return sk

    @classmethod
    def blank(cls) -> "Skeleton":
        prs = pptx.Presentation()
        prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
        sk = cls(prs, False)
        new = lambda: prs.slides.add_slide(prs.slide_layouts[BLANK_LAYOUT])  # noqa: E731
        box = lambda s, g: _textbox(s, Inches(g[0]), Inches(g[1]), Inches(g[2]), Inches(g[3]), [])  # noqa: E731
        sk.cover = new()
        sk.cover_title = box(sk.cover, G_COVER_TITLE)
        _set(box(sk.cover, G_COVER_KICKER), [Line("REFERENCE ARCHITECTURE", 20, HEADER, False, font=FONT_MEDIUM)])
        sk.arch = new()
        sk.arch_title, sk.arch_summary = box(sk.arch, G_TITLE), box(sk.arch, G_SUMMARY)
        sk.goals = [box(sk.arch, (G_GOAL[0], t, G_GOAL[1], G_GOAL[2])) for t in G_GOAL_TOPS]
        sk.cons = new()
        sk.cons_title, sk.cons_sub = box(sk.cons, G_TITLE), box(sk.cons, G_CONS_SUB)
        gt = G_CONS_TABLE
        shape = sk.cons.shapes.add_table(4, 3, Inches(gt[0]), Inches(gt[1]), Inches(gt[2]), Inches(gt[3]))
        sk.cons_table = shape.table
        for j, w in enumerate((2.6, 6.2, 3.0)):
            sk.cons_table.columns[j].width = Inches(w)
        for j, h in enumerate(("Well-Architected Pillar", "Design Decisions & Implementation", "Target Metric Impact")):
            _cell(sk.cons_table.cell(0, j), h, True, HEADER, 11.25)
        sk.app = new()
        sk.app_title = box(sk.app, G_TITLE)
        sk.use_head, sk.avoid_head = box(sk.app, G_HEAD_USE), box(sk.app, G_HEAD_AVOID)
        sk.uses = [box(sk.app, (G_USE_LEFT, t, G_ROW_W, G_ROW_H)) for t in G_ROW_TOPS]
        sk.avoids = [box(sk.app, (G_AVOID_LEFT, t, G_ROW_W, G_ROW_H)) for t in G_ROW_TOPS]
        return sk

    def finish(self) -> None:
        """Drop the template's instruction and example slides (their parts too, so the file shrinks)."""
        if not self.from_template:
            return
        ids = self.prs.slides._sldIdLst
        for idx in sorted(T_DROP, reverse=True):
            if idx < len(ids):
                rid = ids[idx].rId
                del ids[idx]
                try:
                    self.prs.part.drop_rel(rid)
                except (KeyError, AttributeError):
                    pass


def _cell(cell, text: str, bold: bool = False, color: RGBColor = BODY, size: float = 11.25) -> None:
    tf = cell.text_frame
    tf.clear()
    tf.word_wrap = True
    run = tf.paragraphs[0].add_run()
    run.text = _cut(text, MAX_CELL_CHARS)
    run.font.name, run.font.size, run.font.bold = FONT_TABLE, Pt(size), bold
    run.font.color.rgb = color


def _notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text.strip()


# ---------------------------------------------------------------------------------------------- content
def _numbered(i: int, name) -> str:
    name = " ".join(str(name or "").split())
    return name if re.match(r"^\d+[.)]\s", name) else f"{i}. {name}"


def _stage_note(i: int, st: dict) -> str:
    feats = [f["name"] for f in st.get("features", []) if isinstance(f, dict) and f.get("name")]
    note = f"{i}. {st.get('stage')}: {st.get('service')} ({st.get('api')}). {st.get('description', '')}"
    if st.get("model"):
        note += f" Runs on {st['model']} ({st.get('tier')} tier), the newest model verified in this project."
    if feats:
        note += f" Showcases: {', '.join(feats[:3])}."
    if st.get("doc_url"):
        note += f" Source: {st['doc_url']}"
    return note


def _goals(bom: dict, stages: List[dict]) -> List[Tuple[str, str]]:
    """Three design goals: the BOM writer's, else the first three stages' purposes."""
    goals = [(g.get("title", ""), g.get("text", "")) for g in (bom.get("design_goals") or []) if isinstance(g, dict)]
    if len(goals) >= 3:
        return goals[:3]
    return [(_cut(s.get("stage"), 40), _cut(s.get("description"), 120)) for s in stages[:3]]


def _considerations(bom: dict, review: Optional[dict]) -> List[Tuple[str, str, str]]:
    """(pillar, design decision, target metric) x 3: the BOM writer's rows, each followed by the review's finding for
    that pillar when there is one; without a writer, the review alone; without both, what to do."""
    pillars = {p.get("key"): p for p in ((review or {}).get("pillars") or []) if isinstance(p, dict)}
    out = []
    for key, name in (("reliability", "Reliability"), ("cost_optimization", "Cost Optimization"), ("security", "Security")):
        row = (bom.get("considerations") or {}).get(key) if isinstance(bom.get("considerations"), dict) else None
        rev = pillars.get(key)
        if row:
            decision, metric = row.get("decision", ""), row.get("metric", "")
            if rev:
                decision += f" Review: {rev.get('finding', '')}"
                metric += f" · review {rev.get('score')}/5"
        elif rev:
            decision, metric = f"{rev.get('finding', '')} Recommendation: {rev.get('recommendation', '')}", f"{rev.get('score')}/5 in the review"
        else:
            decision, metric = "Written when the demo is rebuilt (the BOM writer did not run for this build).", "-"
        out.append((name, _cut(decision, 230), _cut(metric, 60)))
    return out


def _criteria(bom: dict, key: str) -> List[str]:
    items = [str(x) for x in (bom.get(key) or []) if str(x).strip()]
    if len(items) >= 3:
        return [_cut(x, 120) for x in items[:3]]
    return ["Written when the demo is rebuilt (the BOM writer did not run for this build).", "", ""]


def _cover(sk: Skeleton, customer: str, headline: str, story: dict, ask: str, summary: str) -> None:
    month = _dt.date.today().strftime("%B %Y")
    size = 72 if len(customer) <= 12 else 56 if len(customer) <= 20 else 40
    _set(sk.cover_title, [Line(_cut(customer, 60), size, HEADER, False, font=FONT_MEDIUM),
                          Line(_cut(headline, 90), 24, BLUE, False, font=FONT_MEDIUM),
                          Line(month, 21, MUTED, False, font=FONT_TEXT)])
    _notes(sk.cover, "\n".join(x for x in [
        f"Reference architecture review for {customer}: {headline}.",
        f"The ask: {ask}", f"The solution: {summary}",
        (f"The story: {story.get('logline')}" if story.get("logline") else ""),
        "Agenda: the architecture and its processing sequence; the Well-Architected design considerations; when to "
        "use and when to avoid this architecture."] if x))


def _architecture(sk: Skeleton, customer: str, headline: str, objective: str, goals: List[Tuple[str, str]],
                  stages: List[dict], deliverables: List[dict], story: dict, mode: str) -> None:
    s = sk.arch
    _set(sk.arch_title, [Line(_cut(f"{customer}: {headline}" if headline else f"Reference architecture: {customer}", 95),
                              22, HEADER, True)])
    _set(sk.arch_summary, [Line(_cut(objective, 420), 11, BODY, False, font=FONT_TEXT)])
    for shape, top, (title, text) in zip(sk.goals, G_GOAL_TOPS, goals):
        _place(shape, G_GOAL[0], top, G_GOAL[1], G_GOAL[2])
        _set(shape, [Line(_cut(title, 44), 11, HEADER, True), Line(_cut(text, 130), 10, BODY, False, font=FONT_TEXT)])
    _diagram(s, sk.diagram, stages, deliverables, story)
    ai = [st for st in stages if st.get("model")]
    plain = [st for st in stages if not st.get("model")]
    _notes(s, "\n".join([f"Objective: {objective}", "Design goals: " + "; ".join(f"{t}: {x}" for t, x in goals), "",
                         f"Walk the flow in stage order: {len(stages)} stages, {len(ai)} on Google AI models."]
                        + [_stage_note(i, st) for i, st in enumerate(stages, 1)]
                        + ([f"Demo output: " + "; ".join(f"{d.get('title')} ({manifest.kind_label(d.get('kind'))})"
                                                        for d in deliverables) + "."] if deliverables else [])
                        + ["", "Points to make:",
                           "- Every service, model and feature on this slide cites an official Google doc, read through "
                           "the Developer Knowledge MCP server; nothing here is guessed.",
                           f"- Each AI stage runs on the newest model verified callable in this project ({mode} mode); the "
                           "IDs live in usecase_config.json and are re-verified every 24 hours.",
                           (f"- {len(plain)} stage(s) are Google Cloud services with no model: "
                            + ", ".join(str(st.get("service")) for st in plain) + ".") if plain else
                           "- Every stage runs on a Google AI model.",
                           "- Every shape is editable: move, recolour or annotate it for the customer."]))


def _diagram(s, box: Tuple[float, float, float, float], stages: List[dict], deliverables: List[dict],
             story: dict) -> None:
    """The studio's diagram inside the template's diagram area: stage cards in execution order (one to three rows),
    arrows, a dashed demo-output group at the bottom when the build has outputs, a legend."""
    left, top, width, height = (Inches(v) for v in box)
    n = len(stages)
    if not n:
        _textbox(s, left, top, width, Inches(0.5), [Line("The design lists no stages.", 12, MUTED)])
        return
    legend_h, gap, row_gap = Inches(0.3), Inches(0.2), Inches(0.35)
    out_h = Inches(0.95) if deliverables else 0
    per_row = n if n <= 3 else (n + 1) // 2 if n <= 8 else (n + 2) // 3
    rows = (n + per_row - 1) // per_row
    area_h = height - legend_h - out_h - (Inches(0.15) if deliverables else 0)
    box_w = Emu(int((width - gap * (per_row - 1)) / per_row))
    box_h = Emu(min(int((area_h - row_gap * (rows - 1)) / rows), int(Inches(1.75))))
    size = 8.5 if per_row <= 3 else 8 if per_row == 4 else 7
    positions = []
    for i, st in enumerate(stages):
        r, c = divmod(i, per_row)
        x, y = left + c * (box_w + gap), top + r * (box_h + row_gap)
        positions.append((x, y))
        fill, border = tier_colors(st.get("tier", ""))
        lines = [Line(_cut(_numbered(i + 1, st.get("stage")), 44), size + 0.5, DARK_BLUE, True),
                 Line(_cut(st.get("service"), 36), size, HEADER, True),
                 Line(f"{st['model']}" if st.get("model") else "No AI model", size - 0.5,
                      border if st.get("model") else MUTED, bool(st.get("model")))]
        if box_h >= Inches(1.3):
            lines.append(Line(_cut(st.get("description"), 90 if per_row >= 4 else 120), size - 0.5, BODY))
        _card(s, x, y, box_w, box_h, lines, fill, border, margin=Inches(0.07))
    arrow_w, arrow_h = Inches(0.14), Inches(0.18)
    for i in range(n - 1):
        (x1, y1), (x2, y2) = positions[i], positions[i + 1]
        if y1 == y2:
            _shape(s, MSO_SHAPE.RIGHT_ARROW, x1 + box_w + (gap - arrow_w) // 2, y1 + (box_h - arrow_h) // 2,
                   arrow_w, arrow_h, BLUE)
        else:
            _connector(s, x1 + box_w // 2, y1 + box_h, x2 + box_w // 2, y2, BLUE, dashed=False)
    bottom = top + height
    if deliverables:
        gy = bottom - legend_h - out_h
        _dashed_box(s, left, gy, width, out_h)
        _textbox(s, left + Inches(0.1), gy + Inches(0.04), Inches(2.5), Inches(0.24),
                 [Line("Demo output", 9, HEADER, True)])
        shown, more = deliverables[:MAX_OUTPUT_CARDS], len(deliverables) - MAX_OUTPUT_CARDS
        cw = Emu(int((width - Inches(0.2) - gap * (len(shown) - 1)) / max(len(shown), 1)))
        for j, d in enumerate(shown):
            fill, border = tier_colors(d.get("tier") or "")
            k = len(d.get("variants") or [])
            _card(s, left + Inches(0.1) + j * (cw + gap), gy + Inches(0.3), cw, out_h - Inches(0.38),
                  [Line(_cut(d.get("title"), 40), 8, HEADER, True),
                   Line(f"{manifest.kind_label(d.get('kind'))} · {k} variant{'s' if k != 1 else ''}", 7, MUTED)],
                  fill, border, 0.08, margin=Inches(0.06))
        if more > 0:
            _textbox(s, left + width - Inches(2.3), gy + Inches(0.04), Inches(2.2), Inches(0.24),
                     [Line(f"+ {more} more in the demo", 8, MUTED)])
        lx, ly = positions[-1]
        _connector(s, lx + box_w // 2, ly + box_h, lx + box_w // 2, gy)
    tiers = list(dict.fromkeys(st.get("tier") for st in stages if st.get("model")))
    x, y = left, bottom - legend_h + Inches(0.02)
    for tier in tiers + ([""] if any(not st.get("model") for st in stages) else []):
        fill, border = tier_colors(tier)
        label = TIER_LABELS.get(tier, tier) if tier else "Google Cloud service (no model)"
        sw = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y + Inches(0.04), Inches(0.18), Inches(0.18))
        sw.fill.solid()
        sw.fill.fore_color.rgb = fill if tier else GREY_FILL
        sw.line.color.rgb = border if tier else BORDER
        sw.shadow.inherit = False
        w = Inches(0.25 + 0.07 * len(label))
        _textbox(s, x + Inches(0.22), y, w, Inches(0.26), [Line(label, 8, BODY)])
        x += Inches(0.3) + w
    if story.get("hero"):
        _textbox(s, left, top - Inches(0.27), width, Inches(0.24),
                 [Line(f"For {_cut(story['hero'], 70)}", 8.5, MUTED, False, font=FONT_TEXT)])


def _design_considerations(sk: Skeleton, customer: str, rows: List[Tuple[str, str, str]], review: Optional[dict]) -> None:
    _set(sk.cons_title, [Line("Design considerations", 24, BLUE, True, font=FONT_MEDIUM)])
    done = bool(review) and review.get("status") == "done"
    sub = (f"Well-Architected review: {review.get('verdict')} · average {review.get('average')}/5 (every pillar at or "
           f"above {MIN_PILLAR} means ready for design review)" if done else
           "Design decisions per Well-Architected pillar and the target they serve")
    if sk.cons_sub is not None:
        _set(sk.cons_sub, [Line(_cut(sub, 160), 10.5, MUTED, False, font=FONT_TEXT)])
    table = sk.cons_table
    for i, (pillar, decision, metric) in enumerate(rows, 1):
        if i >= len(table.rows):
            break
        _cell(table.cell(i, 0), pillar, True)
        _cell(table.cell(i, 1), decision)
        _cell(table.cell(i, 2), metric, False, BLUE if i == 1 else GOOGLE_GREEN if i == 2 else RED)
    notes = [f"Design considerations for {customer}, per Well-Architected pillar:"]
    notes += [f"{p}: {d} Target: {m}." for p, d, m in rows]
    if done:
        notes += ["", f"Well-Architected review ({review.get('verdict')}, average {review.get('average')}/5): "
                  + review.get("summary", "")]
        notes += [f"- {p['name']}: {p['score']}/5. {p['finding']} Recommendation: {p['recommendation']} ({p.get('doc_url', '')})"
                  for p in review.get("pillars") or []]
    notes += ["", "Reference: Google Cloud Well-Architected Framework, https://cloud.google.com/architecture/framework"]
    _notes(sk.cons, "\n".join(notes))


def _applicability(sk: Skeleton, customer: str, uses: List[str], avoids: List[str], use_cases: List[str]) -> None:
    _set(sk.app_title, [Line("Applicability criteria", 24, BLUE, True, font=FONT_MEDIUM)])
    _set(sk.use_head, [Line("When to use this reference architecture", 14.5, GOOGLE_GREEN, True, font=FONT_TEXT)])
    _set(sk.avoid_head, [Line("When to avoid this reference architecture", 14.5, RED, True, font=FONT_TEXT)])
    for shapes, icons, texts, x in ((sk.uses, sk.use_icons, uses, G_USE_LEFT), (sk.avoids, sk.avoid_icons, avoids, G_AVOID_LEFT)):
        for shape, icon, text, top in zip(shapes, icons + [None] * 3, texts, G_ROW_TOPS):
            _place(shape, x, top, G_ROW_W, G_ROW_H)
            _set(shape, [Line(text, 11.5, BODY, False, font=FONT_TABLE)])
            if icon is not None:
                icon.left, icon.top = Inches(x - G_ICON_DX), Inches(top + 0.05)
    if sk.from_template:  # the template's right-hand bullets are pushpins: use its red cross instead
        mark = _by_id(sk.app, ID_AVOID_MARK)
        blob = None
        try:
            blob = mark.image.blob if mark is not None else None
        except (AttributeError, ValueError):
            blob = None
        if blob:
            for icon in sk.avoid_icons:
                try:
                    rid = icon._element.xpath(".//a:blip/@r:embed")[0]
                    icon.part.related_part(rid)._blob = blob
                except (IndexError, KeyError, AttributeError):
                    continue
    _notes(sk.app, "\n".join([f"When to use this reference architecture ({customer}):"] + [f"- {u}" for u in uses if u]
                             + ["When to avoid it:"] + [f"- {a}" for a in avoids if a]
                             + (["Use cases and industries where the pattern applies: " + "; ".join(use_cases)] if use_cases
                                else [])))


# ---------------------------------------------------------------------------------------------- entry points
def build_usecase_deck(output_path: str, *, customer: str, ask: str, summary: str, stages: List[dict],
                       rubric: List[dict], attempts: List[dict], files: List[str], whats_new: List[dict], mode: str,
                       deliverables: List[dict], score: float, final_status: str,
                       story: Optional[dict] = None, well_architected: Optional[dict] = None,
                       bom: Optional[dict] = None, template_path: str = "") -> str:
    """Write the four-slide reference architecture deck for one build to `output_path`, on the template at
    `template_path` when it exists (else the blank fallback). -> output_path."""
    story = story if isinstance(story, dict) else {}
    bom = bom if isinstance(bom, dict) and bom.get("status") == "done" else {}
    sk = Skeleton.from_template(template_path) if template_path and os.path.isfile(template_path) else Skeleton.blank()
    headline = bom.get("headline") or ""
    objective = bom.get("objective") or summary
    _cover(sk, customer, headline or "Reference architecture", story, ask, summary)
    _architecture(sk, customer, headline, objective, _goals(bom, stages), stages, deliverables, story, mode)
    _design_considerations(sk, customer, _considerations(bom, well_architected), well_architected)
    _applicability(sk, customer, _criteria(bom, "when_to_use"), _criteria(bom, "when_to_avoid"), bom.get("use_cases") or [])
    sk.finish()
    props = sk.prs.core_properties
    props.title, props.subject = f"{customer}: reference architecture", _cut(ask, 250)
    props.author = props.last_modified_by = "Gemini + MCP Use-Case Studio"
    props.version = DECK_VERSION
    props.category = TEMPLATE_TAG if sk.from_template else BLANK_TAG
    props.keywords = f"{final_status}; rubric {score:.1f}%; {len(attempts)} attempt(s); {len(files)} package files"
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    sk.prs.save(output_path)
    return output_path


def deck_version(path: str) -> str:
    """The DECK_VERSION a saved deck was built with: '' when the file is missing, not a readable deck, or was made
    before decks carried a version (all of which mean: regenerate it)."""
    if not os.path.isfile(path):
        return ""
    try:
        return str(pptx.Presentation(path).core_properties.version or "")
    except Exception:  # a truncated or foreign file: treat as stale, not as an error
        return ""


def deck_template(path: str) -> str:
    """TEMPLATE_TAG when the deck was built on the template, BLANK_TAG for the fallback, '' when unknown."""
    if not os.path.isfile(path):
        return ""
    try:
        return str(pptx.Presentation(path).core_properties.category or "")
    except Exception:
        return ""
