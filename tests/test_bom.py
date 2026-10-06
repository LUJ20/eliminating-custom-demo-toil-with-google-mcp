"""Bill of materials (engine/bom.py): the writer's answer is validated strictly, the four documents render from the
facts with every value escaped, SKILL.md carries the agent-skills frontmatter, the deck carries the narrative, and a
build saved before the BOM existed gets one in place (build_editor.add_bom) unless the switch is off."""
import copy
import dataclasses
import json
import os
import py_compile
import unittest
import zipfile

import pptx

from engine import bom, build_editor, deck_generator as dg, prebuild, versions
from engine import usecase_synthesizer as us
from engine.common import read_json, write_json
from engine.pii_sanitizer import PIISanitizer
from engine.troubleshooter import OutputError
from fakes import OfflineTestCase, seed_registry

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
STAGES = [{"stage": "1. Ingest", "service": "Cloud Storage", "api": "objects.insert", "tier": "", "model": "",
           "location": "", "description": "Stores manuals.", "features": []},
          {"stage": "2. Answer", "service": "Vertex AI", "api": "generateContent", "tier": "fast",
           "model": "gemini-9.9-flash", "location": "global", "description": "Answers questions.", "features": []}]
STORY = {"title": "Mia Fixes The Drill", "logline": "A broken drill is diagnosed from a photo.",
         "hero": "Mia, a site foreman", "challenge": "No manual on site.", "payoff": "Back to work in five minutes.",
         "beats": [{"id": "b1", "title": "The photo", "scene": "Mia sends a photo.", "feature": ""},
                   {"id": "b2", "title": "The answer", "scene": "The agent names the part.", "feature": ""}]}
BLUEPRINT = {"summary": "Answers support questions from manuals.", "stages": STAGES, "story": STORY,
             "deliverables": [{"id": "greeting", "title": "Welcome clip", "kind": "video", "brief": "A greeting.",
                               "beat": "b1", "variants": [{"label": "English", "script": "Welcome to Acme."}]}]}


def good(**changes) -> dict:
    """A complete writer answer for BLUEPRINT, with optional field overrides."""
    data = {
        "headline": "Manual Answering From Photos",
        "objective": "Answer support questions from photos and manuals. Managed services per stage; the model is "
                     "resolved by capability tier.",
        "design_goals": [{"title": f"Goal {i}", "text": f"What goal {i} does."} for i in range(1, 4)],
        "considerations": {k: {"decision": f"{n} is handled by Cloud Run and Cloud Storage.", "metric": f"{n} target"}
                           for k, n in bom.PILLARS},
        "when_to_use": ["Field technicians need answers from manuals.", "Photos are the main input.",
                        "Answers must cite the manual."],
        "when_to_avoid": ["Only structured data is involved.", "No manuals exist.", "A rules engine is enough."],
        "technical_challenges": [{"title": "Long manuals", "text": "Chunked and grounded."},
                                 {"title": "Blurry photos", "text": "Checked before use."}],
        "business_challenges": [{"title": "Downtime", "text": "Minutes matter."},
                                {"title": "Training", "text": "No new tool to learn."}],
        "alternatives": [{"pattern": "Keyword search", "advantages": "Simple.", "disadvantages": "No photos."}],
        "use_cases": ["Field service", "Retail support"],
        "differentiators": ["Vertex AI Search grounds the answers.", "Multimodal input in one call."],
        "segments": [{"title": "Open", "minutes": 2, "beat": "b1", "action": "Show the photo.",
                      "script": "I start with the photo."},
                     {"title": "Answer", "minutes": 0, "beat": "b2", "action": "Play the answer.",
                      "script": "The agent names the part."},
                     {"title": "Close", "minutes": 60, "beat": "b9", "action": "Show the deck.",
                      "script": "I close on the deck."}],
        "poc": {k: [f"{k} practice one.", f"{k} practice two."] for k, _ in bom.POC_PHASES},
        "skills": [{"name": "Manual Answering", "purpose": "Runs stage 2 on the design's model.", "inputs": "A photo.",
                    "outputs": "An answer.", "evaluation": "Judged against the manual."}],
    }
    data.update(changes)
    return data


def narrative(**changes) -> dict:
    return bom.validate(json.dumps(good(**changes)), BLUEPRINT)


class ValidateTest(unittest.TestCase):
    def test_good_answer_is_cleaned_and_clamped(self):
        out = narrative()
        self.assertTrue(bom.done(out))
        self.assertEqual(out["status"], bom.STATUS_DONE)
        self.assertEqual(len(out["design_goals"]), 3)
        self.assertEqual([s["minutes"] for s in out["segments"]], [2, 1, 10])  # clamped to 1-10
        self.assertEqual([s["beat"] for s in out["segments"]], ["b1", "b2", ""])  # only the story's beat ids
        self.assertEqual(out["skills"][0]["name"], "manual-answering")  # slugified
        self.assertEqual(out["differentiators"][0], "Agent Search grounds the answers.")  # current product names
        self.assertEqual(out["considerations"]["security"]["pillar"], "Security")
        self.assertEqual(sorted(out["poc"]), sorted(k for k, _ in bom.POC_PHASES))

    def test_long_text_is_cut_and_extra_entries_dropped(self):
        out = narrative(headline="x" * 100, design_goals=[{"title": f"G{i}", "text": "t"} for i in range(5)])
        self.assertEqual(len(out["headline"]), bom.MAX_HEADLINE)
        self.assertEqual(len(out["design_goals"]), 3)

    def test_the_designs_models_may_be_named_but_no_other(self):
        ok = narrative(when_to_use=["Runs on gemini-9.9-flash.", "Photos.", "Manuals."])
        self.assertIn("gemini-9.9-flash", ok["when_to_use"][0])
        with self.assertRaises(OutputError) as ctx:
            narrative(when_to_use=["Runs on veo-7.0-fake.", "Photos.", "Manuals."])
        self.assertIn("model ID", str(ctx.exception))

    def test_bad_answers_are_rejected(self):
        base = good()
        bad = {
            "not json": "not json at all",
            "a list": "[1, 2]",
            "empty headline": json.dumps(good(headline="")),
            "two goals": json.dumps(good(design_goals=base["design_goals"][:2])),
            "missing pillar": json.dumps(good(considerations={k: v for k, v in base["considerations"].items()
                                                              if k != "security"})),
            "no metric": json.dumps(good(considerations={**base["considerations"],
                                                         "security": {"decision": "d", "metric": ""}})),
            "minutes not a number": json.dumps(good(segments=[dict(s, minutes="soon") for s in base["segments"]])),
            "poc missing": json.dumps(good(poc={})),
            "no skills": json.dumps(good(skills=[])),
        }
        for why, text in bad.items():
            with self.subTest(why), self.assertRaises(OutputError):
                bom.validate(text, BLUEPRINT)

    def test_unavailable_is_not_done(self):
        out = bom.unavailable("the writer timed out")
        self.assertEqual((out["status"], out["reason"]), (bom.STATUS_UNAVAILABLE, "the writer timed out"))
        self.assertFalse(bom.done(out))
        self.assertFalse(bom.done(None))


class DeckTest(OfflineTestCase):
    def test_blank_fallback_deck_carries_the_narrative(self):
        path = dg.build_usecase_deck(os.path.join(self.tmp, "d.pptx"), customer="Acme Tools", ask="ask", summary="sum",
                                     stages=STAGES, rubric=[], attempts=[], files=["pipeline.py"], whats_new=[],
                                     mode="Showcase", deliverables=[], score=90.0, final_status="PASSED", story=STORY,
                                     bom=narrative(), template_path="")
        prs = pptx.Presentation(path)
        self.assertEqual(len(prs.slides), 4)
        self.assertEqual((dg.deck_template(path), dg.deck_version(path)), (dg.BLANK_TAG, dg.DECK_VERSION))
        texts = [" ".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame) for sl in prs.slides]
        self.assertIn("Acme Tools", texts[0])
        self.assertIn("Manual Answering From Photos", texts[0])
        self.assertIn("Goal 1", texts[1])
        table = next(sh for sh in prs.slides[2].shapes if sh.has_table).table
        self.assertIn("Reliability is handled by Cloud Run", table.cell(1, 1).text)
        self.assertEqual(table.cell(1, 2).text, "Reliability target")
        self.assertIn("Field technicians need answers from manuals.", texts[3])
        self.assertIn("Field service", prs.slides[3].notes_slide.notes_text_frame.text)


class DocumentsTest(OfflineTestCase):
    def result(self, bom_value=None) -> dict:
        pd = os.path.join(self.tmp, "acme-tools")
        os.makedirs(pd, exist_ok=True)
        return {"customer_name": "Acme <b>Tools</b>", "usecase_ask": "Support agent.", "summary": "Answers questions.",
                "stages": copy.deepcopy(STAGES), "deliverables": copy.deepcopy(BLUEPRINT["deliverables"]),
                "story": STORY, "build_id": "b1", "whats_new": [], "slug": "acme-tools", "project_dir": pd,
                "eval_metrics": [{"metric": "Code validity", "value": "PASS", "threshold": "compiles", "pass": True,
                                  "notes": ""}],
                "package_files": ["pipeline.py", bom.SKILL_FILE],
                "grounding_sources": [{"title": "Cloud Run", "url": "https://cloud.google.com/run/docs"}],
                "bom": bom_value if bom_value is not None else narrative(headline="Photo <script>alert(1)</script> answers")}

    def test_four_documents_with_every_value_escaped(self):
        docs = bom.render_docs(self.result(), None, "https://docs.google.com/presentation/d/abc/edit")
        self.assertEqual(sorted(docs), sorted(k for k, *_ in bom.DOCS))
        for key, text in docs.items():
            with self.subTest(key):
                self.assertNotIn("<script>", text)
                self.assertNotIn("<b>Tools</b>", text)
                self.assertIn("Acme &lt;b&gt;Tools&lt;/b&gt;", text)
                self.assertIn("Generated by Gemini + MCP Use-Case Studio from build b1", text)
        self.assertIn("&lt;script&gt;", docs["tgd"])
        self.assertIn('href="https://docs.google.com/presentation/d/abc/edit"', docs["tgd"])
        self.assertIn("Agent Platform", docs["tgd"])  # the stage table uses current product names
        self.assertIn("Segment 1: Open", docs["demo"])
        self.assertIn("Welcome to Acme.", docs["demo"])  # the clip script of the segment's beat
        self.assertIn("assessment practice one.", docs["poc"])
        self.assertIn("manual-answering", docs["skills"])
        self.assertIn("name: acme-tools-demo-pipeline", docs["skills"])

    def test_documents_without_a_narrative_say_what_is_missing(self):
        docs = bom.render_docs(self.result(bom.unavailable("the writer timed out")))
        for key, text in docs.items():
            with self.subTest(key):
                self.assertIn("will be written when this demo is rebuilt (the writer timed out)", text)
        self.assertIn("Scene 1: The photo", docs["demo"])  # the story beats stand in for the segments

    def test_write_docs_writes_four_files_once(self):
        res = self.result()
        paths = bom.write_docs(res)
        self.assertEqual(paths, bom.doc_paths(res))
        for key, suffix, _, _ in bom.DOCS:
            self.assertEqual(paths[key], os.path.join(res["project_dir"], bom.FOLDER, f"acme-tools{suffix}"))
            self.assertTrue(os.path.isfile(paths[key]))
        stamps = {k: os.stat(p).st_mtime_ns for k, p in paths.items()}
        self.assertEqual(paths, bom.write_docs(res))  # unchanged content: the files are left untouched
        self.assertEqual(stamps, {k: os.stat(p).st_mtime_ns for k, p in paths.items()})
        self.assertEqual(sorted(os.listdir(os.path.join(res["project_dir"], bom.FOLDER))),
                         sorted(os.path.basename(p) for p in paths.values()))
        res["bom"] = bom.unavailable("the writer timed out")
        bom.write_docs(res)
        with open(paths["tgd"], encoding="utf-8") as f:
            self.assertIn("the writer timed out", f.read())

    def test_titles_and_blurbs(self):
        for key, _, title, blurb in bom.DOCS:
            self.assertEqual((bom.doc_title(key), bom.doc_blurb(key)), (title, blurb))
        self.assertEqual(bom.doc_title("nope"), "nope")


class SkillMarkdownTest(unittest.TestCase):
    RESULT = {"slug": "acme-tools", "customer_name": "Acme Tools", "summary": "Answers questions.",
              "usecase_ask": "Support agent.", "stages": STAGES, "eval_metrics": [],
              "grounding_sources": [{"title": "Cloud Run", "url": "https://cloud.google.com/run/docs"}]}

    def test_frontmatter_and_sections(self):
        md = bom.skill_markdown(self.RESULT, narrative())
        self.assertTrue(md.startswith("---\nname: acme-tools-demo-pipeline\ndescription: "))
        self.assertIn("\n---\n", md[4:])
        for head in ("# Acme Tools demo pipeline", "## When to use this skill", "## Steps", "## The stages",
                     "## Skills in this pipeline", "## Guardrails", "## References"):
            self.assertIn(head, md)
        self.assertIn("**manual-answering**", md)
        self.assertIn("2. **2. Answer** (Agent Platform, gemini-9.9-flash): Answers questions.", md)
        self.assertIn("[Cloud Run](https://cloud.google.com/run/docs)", md)

    def test_without_a_narrative_the_skill_still_runs_the_pipeline(self):
        md = bom.skill_markdown(self.RESULT, bom.unavailable("timed out"))
        self.assertIn("name: acme-tools-demo-pipeline", md)
        self.assertIn("written when the demo is rebuilt", md)
        self.assertIn("Answers questions.", md)


class FakeSynth:
    """Stands in for UseCaseSynthesizer in add_bom: answers write_bom with a canned narrative and counts the calls."""

    def __init__(self, settings, answer: dict):
        self.answer, self.calls, self.seen = answer, 0, None
        self.pii = PIISanitizer.for_settings(settings)

    def write_bom(self, customer: str, ask: str, blueprint: dict) -> dict:
        self.calls += 1
        self.seen = (customer, ask, blueprint)
        return copy.deepcopy(self.answer)


class AddBomTest(OfflineTestCase):
    """Builds saved before the BOM existed get one in place (build_editor.add_bom), honouring BOM_ASSETS."""

    def setUp(self):
        super().setUp()
        seed_registry(self.settings)
        self.settings = dataclasses.replace(self.settings, bom_enabled=True)

    def finished_project(self, bom_value=None) -> str:
        folder = prebuild.project_dir(self.settings, "Acme Tools")
        os.makedirs(folder, exist_ok=True)
        slug = os.path.basename(folder)
        res = {"customer_name": "Acme Tools", "usecase_ask": "Support agent for Acme.", "final_status": "PASSED",
               "mode": self.settings.mode, "score": 100.0, "summary": BLUEPRINT["summary"], "stages": STAGES,
               "deliverables": [], "story": STORY, "models": {}, "generation": us.BUILD_GENERATION, "package_files": [],
               "whats_new": [], "attempt_stats": [], "incidents": [], "grounding_sources": [], "build_id": "b1",
               "eval_metrics": [{"metric": "Code validity", "value": "PASS", "threshold": "compiles",
                                 "method": "programmatic", "notes": "", "pass": True}]}
        if bom_value is not None:
            res["bom"] = bom_value
        write_json(os.path.join(folder, us.RESULT_FILE), res)
        write_json(os.path.join(folder, "usecase_config.json"), {"summary": res["summary"], "stages": res["stages"]})
        for name, text in (("pipeline.py", "print('hi')\n"), ("requirements.txt", ""), ("README.md", "# x\n")):
            with open(os.path.join(folder, name), "w", encoding="utf-8") as f:
                f.write(text)
        with zipfile.ZipFile(os.path.join(folder, f"{slug}_codebase.zip"), "w") as zf:
            zf.writestr("pipeline.py", "print('hi')\n")
        return folder

    def add(self, folder: str, answer=None, settings=None):
        settings = settings or self.settings
        res = build_editor.load_result(settings, folder)
        synth = FakeSynth(settings, answer if answer is not None else narrative())
        return build_editor.add_bom(settings, res, synth), res, synth

    def test_a_saved_build_gets_its_bom_in_place_and_stays_clean(self):
        folder = self.finished_project()
        slug = os.path.basename(folder)
        added, res, synth = self.add(folder)
        self.assertTrue(added)
        self.assertEqual(synth.calls, 1)
        self.assertEqual(synth.seen[0], "Acme Tools")
        self.assertEqual(synth.seen[2]["story"]["title"], STORY["title"])
        self.assertTrue(bom.done(res["bom"]))
        self.assertIn(bom.SKILL_FILE, res["package_files"])
        stored = read_json(os.path.join(folder, us.RESULT_FILE), {})
        self.assertEqual((stored["bom"]["status"], stored["bom"]["headline"]), ("done", "Manual Answering From Photos"))
        self.assertEqual(stored["final_status"], "PASSED")
        with open(os.path.join(folder, bom.SKILL_FILE), encoding="utf-8") as f:
            text = f.read()
        self.assertIn("-demo-pipeline", text)
        self.assertIn("# Acme Tools demo pipeline", text)
        with zipfile.ZipFile(os.path.join(folder, f"{slug}_codebase.zip")) as zf:
            self.assertIn(bom.SKILL_FILE, zf.namelist())
            self.assertIn("pipeline.py", zf.namelist())
        deck = os.path.join(folder, f"{slug}_architecture_deck.pptx")
        self.assertEqual(dg.deck_version(deck), dg.DECK_VERSION)
        prs = pptx.Presentation(deck)
        self.assertEqual(len(prs.slides), 4)
        cover = " ".join(sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame)
        self.assertIn("Manual Answering From Photos", cover)
        self.assertFalse(versions.is_dirty(folder))  # recorded as the saved version: no "Unsaved changes" row
        self.assertEqual(build_editor.load_result(self.settings, folder)["bom"]["headline"], "Manual Answering From Photos")
        added, _, synth = self.add(folder)  # nothing left to do
        self.assertEqual((added, synth.calls), (False, 0))

    def test_an_unavailable_narrative_changes_nothing_and_is_tried_again(self):
        folder = self.finished_project(bom.unavailable("timed out"))
        added, _, synth = self.add(folder, answer=bom.unavailable("still down"))
        self.assertEqual((added, synth.calls), (False, 1))
        self.assertFalse(os.path.exists(os.path.join(folder, bom.SKILL_FILE)))
        self.assertEqual(read_json(os.path.join(folder, us.RESULT_FILE), {})["bom"]["reason"], "timed out")
        self.assertFalse(versions.is_dirty(folder))
        added, _, _ = self.add(folder)  # the next attempt, with a good answer, succeeds
        self.assertTrue(added)

    def test_a_project_with_unsaved_edits_is_left_alone(self):
        folder = self.finished_project()
        versions.ensure_baseline(folder)
        with open(os.path.join(folder, "pipeline.py"), "a", encoding="utf-8") as f:
            f.write("print('edited in chat')\n")
        self.assertTrue(versions.is_dirty(folder))
        added, _, synth = self.add(folder)
        self.assertEqual((added, synth.calls), (False, 0))
        self.assertFalse(os.path.exists(os.path.join(folder, bom.SKILL_FILE)))

    def test_the_switch_turns_the_backfill_off(self):
        folder = self.finished_project()
        added, _, synth = self.add(folder, settings=dataclasses.replace(self.settings, bom_enabled=False))
        self.assertEqual((added, synth.calls), (False, 0))
        self.assertFalse(os.path.exists(os.path.join(folder, bom.SKILL_FILE)))


class AppCompilesTest(unittest.TestCase):
    def test_app_compiles(self):
        py_compile.compile(os.path.join(ROOT, "app.py"), doraise=True)


if __name__ == "__main__":
    unittest.main()
