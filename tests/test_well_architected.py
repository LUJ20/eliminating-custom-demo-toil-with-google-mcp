"""Well-Architected review (engine/well_architected.py): Framework pages from MCP (Framework pages first, no
duplicates, capped, cached), strict validation of the reviewer's answer, the readiness gate, the package markdown,
the deck slide, and the synthesizer integration (advisory: never changes the score; off switch; MCP outage shows
'Not reviewed')."""
import copy
import dataclasses
import json
import os
import unittest
import zipfile
from unittest import mock

import pptx

from engine import acceptance
from engine import build_editor
from engine import deck_generator as dg
from engine import deliverables as dlv
from engine import usecase_synthesizer as us
from engine import versions
from engine import well_architected as waf
from engine.common import read_json, write_json
from engine.deck_generator import DECK_VERSION, deck_version
from engine.troubleshooter import OutputError
from fakes import FakeMcp, OfflineTestCase, seed_registry

_SYNTH = us.UseCaseSynthesizer  # the real class, for tests that patch the module attribute

FW = "documents/docs.cloud.google.com/architecture/framework"
HITS = [
    {"parent": "documents/docs.cloud.google.com/run/docs/configuring/cpu", "content": "Cloud Run CPU allocation."},
    {"parent": f"{FW}/security", "content": "Security, privacy and compliance pillar: least privilege, CMEK."},
    {"parent": f"{FW}/reliability", "content": "Reliability pillar: redundancy, graceful degradation."},
    {"parent": f"{FW}/security", "content": "duplicate of the security page"},
    {"parent": f"{FW}/reliability/printable", "content": "print copy of the reliability page"},
]
BLUEPRINT = {"summary": "Answers support questions from manuals.", "deliverables": [{"kind": "text"}], "story": {},
             "stages": [{"stage": "1. Ingest", "service": "Cloud Storage", "api": "objects.insert", "tier": "",
                         "model": "", "location": "", "description": "Stores manuals.", "features": []},
                        {"stage": "2. Answer", "service": "Vertex AI", "api": "generateContent", "tier": "fast",
                         "model": "flash-x", "location": "global", "description": "Answers questions.", "features": []}]}


def answer(scores=(4, 5, 3, 4, 4), **changes):
    """A reviewer answer with these pillar scores (in PILLARS order) and optional per-pillar overrides."""
    pillars = {k: {"score": s, "finding": f"{n} finding.", "recommendation": f"{n} fix.", "source": 1}
               for (k, n, _), s in zip(waf.PILLARS, scores)}
    for k, fields in changes.items():
        pillars[k].update(fields)
    return json.dumps({"pillars": pillars, "summary": "Solid demo design. Add IAM scoping before the pilot."})


def docs(n=3):
    return [{"parent": f"{FW}/p{i}", "title": f"Page {i}", "url": f"https://docs.cloud.google.com/architecture/framework/p{i}",
             "snippet": "..."} for i in range(1, n + 1)]


class FrameworkDocsTest(unittest.TestCase):
    def setUp(self):
        waf.reset_cache()
        self.addCleanup(waf.reset_cache)

    def test_framework_pages_first_deduplicated_and_one_query_per_pillar(self):
        mcp = FakeMcp(results=HITS)
        got = waf.framework_docs(mcp)
        self.assertEqual([d["url"] for d in got],
                         ["https://docs.cloud.google.com/architecture/framework/security",
                          "https://docs.cloud.google.com/architecture/framework/reliability",
                          "https://docs.cloud.google.com/run/docs/configuring/cpu"])
        self.assertEqual(got[0]["snippet"], "Security, privacy and compliance pillar: least privilege, CMEK.")
        self.assertTrue(got[0]["title"])
        self.assertEqual(len(mcp.queries), len(waf.PILLARS) + 1)
        self.assertEqual(set(mcp.queries), {q for _, _, q in waf.PILLARS} | {waf.PERSPECTIVE_QUERY})

    def test_capped_at_max_docs(self):
        mcp = FakeMcp(results=[{"parent": f"{FW}/page{i}", "content": "x"} for i in range(waf.MAX_DOCS + 5)])
        self.assertEqual(len(waf.framework_docs(mcp)), waf.MAX_DOCS)

    def test_cached_for_a_day_and_an_empty_answer_is_not_cached(self):
        empty = FakeMcp()
        self.assertEqual(waf.framework_docs(empty, now=1000.0), [])
        mcp = FakeMcp(results=HITS)
        first = waf.framework_docs(mcp, now=1000.0)
        self.assertEqual(len(mcp.queries), len(waf.PILLARS) + 1)
        self.assertEqual(waf.framework_docs(mcp, now=1000.0 + waf.DOCS_TTL_S - 1), first)
        self.assertEqual(len(mcp.queries), len(waf.PILLARS) + 1)  # served from the cache
        waf.framework_docs(mcp, now=1000.0 + waf.DOCS_TTL_S + 1)
        self.assertEqual(len(mcp.queries), 2 * (len(waf.PILLARS) + 1))  # looked up again after the TTL

    def test_results_without_a_document_parent_are_skipped(self):
        self.assertEqual(waf.framework_docs(FakeMcp(results=[{"parent": "", "content": "x"}, {"content": "y"}])), [])


class ValidateReviewTest(unittest.TestCase):
    def test_good_answer_is_scored_in_pillar_order_with_its_sources(self):
        rev = waf.validate_review(answer(cost_optimization={"source": 3}), docs())
        self.assertEqual([p["key"] for p in rev["pillars"]], [k for k, _, _ in waf.PILLARS])
        self.assertEqual([p["score"] for p in rev["pillars"]], [4, 5, 3, 4, 4])
        self.assertEqual((rev["average"], rev["ready"], rev["verdict"], rev["status"]), (4.0, True, waf.READY, "done"))
        self.assertEqual(rev["pillars"][3]["doc_url"], "https://docs.cloud.google.com/architecture/framework/p3")
        self.assertEqual(rev["pillars"][0]["doc_title"], "Page 1")
        self.assertEqual(rev["pillars"][1]["name"], "Security, privacy and compliance")
        self.assertTrue(rev["summary"].startswith("Solid demo design."))

    def test_one_weak_pillar_means_needs_work(self):
        rev = waf.validate_review(answer(scores=(5, 2, 5, 5, 5)), docs())
        self.assertEqual((rev["ready"], rev["verdict"], rev["average"]), (False, waf.NEEDS_WORK, 4.4))

    def test_scores_at_the_threshold_are_ready(self):
        self.assertTrue(waf.validate_review(answer(scores=(3, 3, 3, 3, 3)), docs())["ready"])

    def test_bad_answers_are_rejected_with_a_reason(self):
        cases = [
            ("not json", "valid JSON"),
            (json.dumps({"summary": "x"}), '"pillars" object'),
            (json.dumps({"pillars": {"security": {}}, "summary": "x"}), "missing the pillar operational_excellence"),
            (answer(security={"score": 6}), "security score must be 1-5"),
            (answer(security={"score": "high"}), "security needs a 1-5 score"),
            (answer(reliability={"finding": ""}), "reliability needs a finding"),
            (answer(reliability={"recommendation": "  "}), "reliability needs a finding"),
            (answer(cost_optimization={"source": 4}), "cites source 4, but only 1-3 exist"),
            (answer(cost_optimization={"source": 0}), "cites source 0"),
            (answer(cost_optimization={"source": None}), "needs a source number"),
            (json.dumps({"pillars": json.loads(answer())["pillars"], "summary": ""}), '"summary"'),
        ]
        for text, reason in cases:
            with self.subTest(reason=reason), self.assertRaises(OutputError) as cm:
                waf.validate_review(text, docs())
            self.assertIn(reason, str(cm.exception))

    def test_fractional_scores_are_rounded_and_long_text_is_cut(self):
        rev = waf.validate_review(answer(security={"score": 4.6, "finding": "word " * 200}), docs())
        self.assertEqual(rev["pillars"][1]["score"], 5)
        self.assertLessEqual(len(rev["pillars"][1]["finding"]), waf.MAX_TEXT)

    def test_review_needs_docs_and_passes_the_hint_to_the_model(self):
        settings = mock.Mock()
        with self.assertRaises(OutputError):
            waf.review(settings, "m", "global", "", customer="Acme", ask="a", blueprint=BLUEPRINT, docs=[])
        with mock.patch("engine.vertex.generate", return_value=(answer(), 1)) as gen:
            rev, quality = waf.review(settings, "m", "global", "score must be 1-5", customer="Acme", ask="a",
                                      blueprint=BLUEPRINT, docs=docs())
        self.assertEqual((rev["status"], quality), ("done", 1.0))
        prompt = gen.call_args.args[2]
        self.assertIn("Your previous answer was rejected: score must be 1-5", prompt)
        self.assertIn("[3] Page 3 (https://docs.cloud.google.com/architecture/framework/p3)", prompt)
        self.assertIn('"service": "Cloud Storage"', prompt)
        self.assertTrue(gen.call_args.kwargs["json_mode"])
        self.assertEqual(gen.call_args.kwargs["location"], "global")

    def test_unavailable(self):
        rev = waf.unavailable("MCP down")
        self.assertEqual((rev["status"], rev["verdict"], rev["ready"], rev["pillars"]), ("unavailable", "Not reviewed", False, []))
        self.assertEqual(rev["summary"], "MCP down")


class MarkdownAndDeckTest(unittest.TestCase):
    def test_markdown_has_the_verdict_and_one_row_per_pillar_with_its_page(self):
        rev = waf.validate_review(answer(scores=(5, 2, 5, 5, 5)), docs())
        md = waf.to_markdown(rev, "Acme")
        self.assertTrue(md.startswith("# Well-Architected review: Acme"))
        self.assertIn(f"**Verdict: {waf.NEEDS_WORK}** (average 4.4/5", md)
        self.assertEqual(md.count("| Security, privacy and compliance | 2/5 |"), 1)
        self.assertEqual(md.count("[Page 1](https://docs.cloud.google.com/architecture/framework/p1)"), 5)
        self.assertIn("not a certification", md)

    def test_markdown_for_a_review_that_did_not_run(self):
        self.assertIn("**Not reviewed**: MCP down", waf.to_markdown(waf.unavailable("MCP down"), "Acme"))
        self.assertIn("**Not reviewed**: the review did not run for this build", waf.to_markdown(None, "Acme"))

    def test_deck_gets_a_review_slide_only_when_the_build_has_one(self):
        tmp = os.path.join(os.path.dirname(__file__), ".tmp_waf_deck.pptx")
        self.addCleanup(lambda: os.path.exists(tmp) and os.remove(tmp))
        rev = waf.validate_review(answer(scores=(5, 2, 5, 5, 5)), docs())
        for review, slides in ((None, 5), (rev, 6), (waf.unavailable("MCP down"), 6)):
            with self.subTest(review=(review or {}).get("status")):
                dg.build_usecase_deck(tmp, customer="Acme", ask="a", summary="s", stages=BLUEPRINT["stages"], rubric=[],
                                      attempts=[], files=["pipeline.py"], whats_new=[], mode="live", deliverables=[],
                                      score=90.0, final_status="PASSED", well_architected=review)
                prs = pptx.Presentation(tmp)
                self.assertEqual(len(prs.slides), slides)
                if review:
                    texts = "\n".join(sh.text_frame.text for sh in prs.slides[4].shapes if sh.has_text_frame)
                    self.assertIn("Well-Architected review for Acme", texts)
                    self.assertIn(review["verdict"], texts)
                    if review["status"] == "done":
                        table = next(sh.table for sh in prs.slides[4].shapes if sh.has_table)
                        self.assertEqual(len(table.rows), 1 + len(waf.PILLARS))
                        self.assertEqual(table.cell(2, 1).text, "2/5")


class SynthesizerIntegrationTest(OfflineTestCase):
    def setUp(self):
        super().setUp()
        waf.reset_cache()
        self.addCleanup(waf.reset_cache)

    def build(self, enabled=True, mcp=None, generate=None):
        seed_registry(self.settings)
        synth = us.UseCaseSynthesizer(dataclasses.replace(self.settings, well_architected_enabled=enabled),
                                      mcp=mcp if mcp is not None else FakeMcp(results=HITS))
        files = {"pipeline.py": "print('hi')\n", "requirements.txt": "", "README.md": "# x\n", "usecase_config.json": "{}"}
        rubric = [{"metric": "Code validity", "value": "PASS", "threshold": "compiles", "notes": "", "method": "programmatic",
                   "pass": True}]
        a = us.Attempt(copy.deepcopy(BLUEPRINT), files, {"redactions_applied": []}, rubric, 100.0, True, True, [],
                       {"planner": "p", "coder": "c", "judge": "j"})

        def fake_attempt(n, customer, ask, grounding, feedback, say, frozen=None, on_plan=None):
            on_plan(a.blueprint)
            return a

        with mock.patch.object(synth, "_prepare_models"), mock.patch.object(synth, "_whats_new", return_value=[]), \
                mock.patch.object(synth, "_write_outputs") as write, mock.patch.object(synth, "_grounding", return_value=[]), \
                mock.patch.object(synth, "_attempt", side_effect=fake_attempt), mock.patch.object(dlv, "start"), \
                mock.patch("engine.vertex.generate", side_effect=generate or (lambda *a, **k: (answer(scores=(5, 2, 5, 5, 5)), 1))):
            result = synth.resolve_and_build("Acme", "Support agent for Acme: answer from manuals.")
        return result, write, synth

    def test_review_ships_in_the_result_and_the_package_without_touching_the_score(self):
        result, write, synth = self.build()
        rev = result["well_architected"]
        self.assertEqual((rev["status"], rev["verdict"], rev["average"]), ("done", waf.NEEDS_WORK, 4.4))
        self.assertEqual(rev["model"], "gemini-3.1-pro-preview")  # the reasoning tier, picked by the resolver
        self.assertEqual(rev["pillars"][0]["doc_url"], "https://docs.cloud.google.com/architecture/framework/security")
        self.assertEqual((result["final_status"], result["score"]), ("PASSED", 100.0))  # advisory
        self.assertNotIn(acceptance.ACCEPTANCE_METRIC, [r["metric"] for r in result["eval_metrics"]])
        files = write.call_args.args[1]
        self.assertIn(waf.REVIEW_FILE, files)
        self.assertIn(f"**Verdict: {waf.NEEDS_WORK}**", files[waf.REVIEW_FILE])
        self.assertIn("well_architected", json.loads(us.persisted(result)))
        self.assertEqual(synth.doctor.incidents, [])

    def test_disabled_means_no_lookup_no_model_call_and_no_key(self):
        mcp = FakeMcp(results=HITS)
        gen = mock.Mock(side_effect=AssertionError("the review must not call a model when it is off"))
        result, write, _ = self.build(enabled=False, mcp=mcp, generate=gen)
        self.assertNotIn("well_architected", result)
        self.assertEqual(mcp.queries, [])
        self.assertNotIn(waf.REVIEW_FILE, write.call_args.args[1])

    def test_no_framework_pages_means_not_reviewed_never_a_failed_build(self):
        gen = mock.Mock(side_effect=AssertionError("no review without Framework pages"))
        result, write, _ = self.build(mcp=FakeMcp(), generate=gen)
        rev = result["well_architected"]
        self.assertEqual((rev["status"], rev["verdict"]), ("unavailable", "Not reviewed"))
        self.assertIn("could not be retrieved", rev["summary"])
        self.assertEqual((result["final_status"], result["score"]), ("PASSED", 100.0))
        self.assertIn("**Not reviewed**", write.call_args.args[1][waf.REVIEW_FILE])

    def test_a_bad_answer_is_re_prompted_with_the_reason(self):
        calls = []

        def generate(settings, model, prompt, **kw):
            calls.append(prompt)
            return (answer(security={"score": 9}) if len(calls) == 1 else answer(), 1)

        result, _, _ = self.build(generate=generate)
        self.assertEqual(result["well_architected"]["verdict"], waf.READY)
        self.assertEqual(len(calls), 2)
        self.assertIn("rejected: pillar security score must be 1-5", calls[1])


class SavedBuildBackfillTest(OfflineTestCase):
    """Builds saved before the review existed get one in place (build_editor.add_review via prebuild.refresh_reviews)."""

    def setUp(self):
        super().setUp()
        seed_registry(self.settings)
        self.settings = dataclasses.replace(self.settings, well_architected_enabled=True)
        waf.reset_cache()
        self.addCleanup(waf.reset_cache)

    def finished_project(self, review=None) -> str:
        from engine import prebuild
        folder = prebuild.project_dir(self.settings, "Acme Tools")
        os.makedirs(folder, exist_ok=True)
        slug = os.path.basename(folder)
        res = {"customer_name": "Acme Tools", "usecase_ask": "Support agent for Acme.", "final_status": "PASSED",
               "mode": self.settings.mode, "score": 100.0, "summary": BLUEPRINT["summary"], "stages": BLUEPRINT["stages"],
               "deliverables": [], "story": {}, "models": {}, "generation": us.BUILD_GENERATION, "package_files": [],
               "whats_new": [], "attempt_stats": [], "incidents": [], "grounding_sources": [], "build_id": "b1",
               "eval_metrics": [{"metric": "Code validity", "value": "PASS", "threshold": "compiles", "method": "programmatic",
                                 "notes": "", "pass": True}]}
        if review is not None:
            res["well_architected"] = review
        write_json(os.path.join(folder, us.RESULT_FILE), res)
        write_json(os.path.join(folder, "usecase_config.json"), {"summary": res["summary"], "stages": res["stages"]})
        for name, text in (("pipeline.py", "print('hi')\n"), ("requirements.txt", ""), ("README.md", "# x\n")):
            with open(os.path.join(folder, name), "w", encoding="utf-8") as f:
                f.write(text)
        with zipfile.ZipFile(os.path.join(folder, f"{slug}_codebase.zip"), "w") as zf:
            zf.writestr("pipeline.py", "print('hi')\n")
        return folder

    def sweep(self, framework=None, generate=None):
        from engine import prebuild
        with mock.patch.object(waf, "framework_docs", return_value=docs() if framework is None else framework), \
                mock.patch.object(us, "UseCaseSynthesizer", lambda s: _SYNTH(s, mcp=FakeMcp())), \
                mock.patch("engine.vertex.generate", side_effect=generate or (lambda *a, **k: (answer(), 1))) as gen:
            done = prebuild.refresh_reviews(self.settings, log=lambda m: None)
        return done, gen

    def test_a_saved_build_gets_its_review_in_place_and_stays_clean(self):
        folder = self.finished_project()
        slug = os.path.basename(folder)
        done, gen = self.sweep()
        self.assertEqual(done, [slug])
        self.assertEqual(gen.call_count, 1)
        stored = read_json(os.path.join(folder, us.RESULT_FILE), {})
        self.assertEqual((stored["well_architected"]["status"], stored["well_architected"]["verdict"]), ("done", waf.READY))
        self.assertEqual(stored["final_status"], "PASSED")
        with open(os.path.join(folder, waf.REVIEW_FILE), encoding="utf-8") as f:
            self.assertIn(f"**Verdict: {waf.READY}**", f.read())
        with zipfile.ZipFile(os.path.join(folder, f"{slug}_codebase.zip")) as zf:
            self.assertIn(waf.REVIEW_FILE, zf.namelist())
            self.assertIn("pipeline.py", zf.namelist())
        deck = os.path.join(folder, f"{slug}_architecture_deck.pptx")
        self.assertEqual(deck_version(deck), DECK_VERSION)
        self.assertEqual(len(pptx.Presentation(deck).slides), 6)
        self.assertFalse(versions.is_dirty(folder))  # recorded as the saved version: no "Unsaved changes" row
        loaded = build_editor.load_result(self.settings, folder)
        self.assertIn(waf.REVIEW_FILE, loaded["package_files"])
        self.assertEqual(loaded["well_architected"]["average"], 4.0)
        done, gen = self.sweep()  # a second sweep has nothing to do
        self.assertEqual((done, gen.call_count), ([], 0))

    def test_a_review_that_could_not_run_before_is_tried_again(self):
        folder = self.finished_project(review=waf.unavailable("MCP was down"))
        done, _ = self.sweep()
        self.assertEqual(done, [os.path.basename(folder)])
        self.assertEqual(read_json(os.path.join(folder, us.RESULT_FILE), {})["well_architected"]["status"], "done")

    def test_no_framework_pages_leaves_the_project_untouched(self):
        folder = self.finished_project()
        gen = mock.Mock(side_effect=AssertionError("no review without Framework pages"))
        done, _ = self.sweep(framework=[], generate=gen)
        self.assertEqual(done, [])
        self.assertFalse(os.path.exists(os.path.join(folder, waf.REVIEW_FILE)))
        self.assertNotIn("well_architected", read_json(os.path.join(folder, us.RESULT_FILE), {}))
        self.assertFalse(versions.is_dirty(folder))

    def test_a_project_with_unsaved_edits_is_left_alone(self):
        folder = self.finished_project()
        versions.ensure_baseline(folder)
        with open(os.path.join(folder, "pipeline.py"), "a", encoding="utf-8") as f:
            f.write("print('edited in chat')\n")
        self.assertTrue(versions.is_dirty(folder))
        done, gen = self.sweep()
        self.assertEqual((done, gen.call_count), ([], 0))
        self.assertFalse(os.path.exists(os.path.join(folder, waf.REVIEW_FILE)))

    def test_the_switch_turns_the_sweep_off(self):
        from engine import prebuild
        self.finished_project()
        with mock.patch.object(us, "UseCaseSynthesizer", side_effect=AssertionError("no synthesizer when off")):
            self.assertEqual(prebuild.refresh_reviews(dataclasses.replace(self.settings, well_architected_enabled=False)), [])

    def test_start_background_reviews_after_the_builds_and_before_the_chats(self):
        from engine import prebuild
        order = []
        with mock.patch.object(prebuild, "refresh_decks", side_effect=lambda *a, **k: order.append("decks") or []), \
                mock.patch.object(prebuild, "run", side_effect=lambda *a, **k: order.append("run") or {"built": []}), \
                mock.patch.object(prebuild, "refresh_reviews", side_effect=lambda *a, **k: order.append("reviews") or []), \
                mock.patch.object(prebuild, "refresh_chats", side_effect=lambda *a, **k: order.append("chats") or []), \
                mock.patch.object(prebuild, "POLL_S", 0):
            th = prebuild.start_background(self.settings)
            th.join(5)
        self.assertEqual(order, ["decks", "run", "reviews", "chats"])

    def test_reviews_only_cli_does_not_build(self):
        from engine import prebuild
        with mock.patch.object(prebuild, "get_settings", return_value=self.settings), \
                mock.patch.object(prebuild, "refresh_reviews", return_value=["acme_tools"]) as sweep, \
                mock.patch.object(prebuild, "build") as build, mock.patch("builtins.print") as out:
            self.assertEqual(0, prebuild.main(["--reviews"]))
        build.assert_not_called()
        sweep.assert_called_once()
        self.assertIn("1 project(s) reviewed", "\n".join(str(c.args[0]) for c in out.call_args_list))


if __name__ == "__main__":
    unittest.main()
