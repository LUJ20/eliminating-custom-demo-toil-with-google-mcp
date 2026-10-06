"""Storytelling: the planner's story is validated and threaded to every deliverable, the media director writes
scripts from it, the deck opens with it, and builds made before stories keep their clips."""
import copy
import os
import unittest
from unittest import mock

import pptx

from engine import brain, build_editor, deliverables as dlv, manifest, usecase_synthesizer as us, vertex
from engine.deck_generator import build_usecase_deck
from engine.troubleshooter import OutputError
from fakes import OfflineTestCase
from test_deliverables import CATALOG, STORY, concierge_manifest


def items(beat="b2"):
    out = manifest.validate_manifest([dict(d, beat=beat) for d in concierge_manifest()], CATALOG, 16)
    manifest.attach_models(out, CATALOG)
    return out


class ValidateStoryTest(unittest.TestCase):
    def test_story_threads_scene_text_to_deliverables(self):
        ds = items()
        story = brain.validate_story(copy.deepcopy(STORY), ds)
        self.assertEqual([b["id"] for b in story["beats"]], ["b1", "b2", "b3"])
        self.assertTrue(all(d["scene"] == "The concierge answers in Japanese." for d in ds))

    def test_problems_the_planner_must_fix(self):
        bad = {
            "no story": None,
            "missing hero": {**STORY, "hero": ""},
            "too few beats": {**STORY, "beats": STORY["beats"][:1]},
            "beat without scene": {**STORY, "beats": [dict(b, scene="") for b in STORY["beats"]]},
            "model id": {**STORY, "payoff": "served by gemini-3.1-flash"},
        }
        for why, raw in bad.items():
            with self.subTest(why), self.assertRaises(OutputError):
                brain.validate_story(copy.deepcopy(raw), items())
        with self.subTest("unknown beat"), self.assertRaises(OutputError):
            brain.validate_story(copy.deepcopy(STORY), items(beat="b9"))
        with self.subTest("no deliverable plays a scene"), self.assertRaises(OutputError):
            brain.validate_story(copy.deepcopy(STORY), items(beat=""))

    def test_judge_scores_storytelling(self):
        self.assertIn("storytelling", brain.CRITERIA)
        self.assertIn("storytelling", brain.CRITERIA_HELP)


class DirectorStoryTest(OfflineTestCase):
    def test_director_prompt_carries_the_story_and_scene(self):
        ds = items()
        brain.validate_story(copy.deepcopy(STORY), ds)
        seen = {}

        def fake_generate(settings, model, prompt, **kw):
            seen["prompt"] = prompt
            raise vertex.VertexError(500, "stop here", model)

        with mock.patch.object(vertex, "generate", side_effect=fake_generate), self.assertRaises(vertex.VertexError):
            brain.direct_media(self.settings, "m", "global", "", customer="Cymbal Air", ask="a", summary="s",
                               deliverables=ds, story=STORY)
        self.assertIn("Aiko Makes Her Connection", seen["prompt"])
        self.assertIn("The concierge answers in Japanese.", seen["prompt"])
        self.assertIn("never a generic greeting", seen["prompt"])

    def test_no_story_block_without_story(self):
        self.assertEqual(brain.story_block(None), "")
        self.assertEqual(brain.story_block({}), "")


class DeliverablesStoryTest(OfflineTestCase):
    def test_status_keeps_the_story_and_later_starts_inherit_it(self):
        pd = os.path.join(self.tmp, "p")
        ds = items()
        brain.validate_story(copy.deepcopy(STORY), ds)
        with mock.patch.object(dlv._Job, "run", lambda self: None):
            dlv.start(self.settings, build_id="b1", project_dir=pd, customer="c", ask="a", summary="s",
                      deliverables=ds, story=STORY)
            self.assertEqual(dlv.load_status(pd)["story"]["title"], STORY["title"])
            self.assertEqual(dlv.load_status(pd)["deliverables"][1]["scene"], "The concierge answers in Japanese.")
            dlv.start(self.settings, build_id="b2", project_dir=pd, customer="c", ask="a", summary="s",
                      deliverables=ds)  # an edit that does not pass the story keeps it
            self.assertEqual(dlv.load_status(pd)["story"]["title"], STORY["title"])

    def test_builds_without_a_story_keep_their_hashes(self):
        plain = manifest.validate_manifest(concierge_manifest(), CATALOG, 16)
        legacy = [{k: v for k, v in d.items() if k != "beat"} for d in plain]
        self.assertEqual(dlv.spec_hashes(plain), dlv.spec_hashes(legacy))

    def test_a_changed_scene_changes_the_hash(self):
        a, b = items(), items()
        brain.validate_story(copy.deepcopy(STORY), a)
        other = copy.deepcopy(STORY)
        other["beats"][1]["scene"] = "The concierge rebooks her in Japanese."
        brain.validate_story(other, b)
        self.assertNotEqual(dlv.spec_hashes(a)["greeting"], dlv.spec_hashes(b)["greeting"])


class DeckAndPackageStoryTest(OfflineTestCase):
    def test_deck_opens_with_the_story(self):
        ds = items()
        path = build_usecase_deck(os.path.join(self.tmp, "d.pptx"), customer="Cymbal Air", ask="ask", summary="sum",
                                  stages=[], rubric=[], attempts=[], files=["pipeline.py"], whats_new=[],
                                  mode="Showcase", deliverables=ds, score=90.0, final_status="PASSED", story=STORY)
        prs = pptx.Presentation(path)
        self.assertEqual(len(prs.slides), 5)
        first = " ".join(sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame)
        self.assertIn("The story: Aiko Makes Her Connection", first)
        self.assertIn("Warm handoff", first)
        self.assertIn("Rebooked in 3 minutes", first)
        self.assertIn("Close: Rebooked", prs.slides[0].notes_slide.notes_text_frame.text)

    def test_story_accepted_needs_coverage_and_story(self):
        cov, story = brain.CRITERIA["deliverable_coverage"], brain.CRITERIA["storytelling"]
        self.assertTrue(us._story_accepted([{"metric": cov, "pass": True}, {"metric": story, "pass": True}]))
        self.assertFalse(us._story_accepted([{"metric": cov, "pass": True}, {"metric": story, "pass": False}]))
        self.assertTrue(us._story_accepted([{"metric": cov, "pass": True}]))  # judged before stories existed

    def test_readme_section(self):
        md = us._story_md(STORY)
        self.assertIn("## The demo story: Aiko Makes Her Connection", md)
        self.assertIn("3. **Warm handoff**", md)
        self.assertEqual(us._story_md({}), "")


class EditorStoryTest(unittest.TestCase):
    def test_chat_edits_stay_in_the_story_and_keep_hashes(self):
        current = items()
        story = brain.validate_story(copy.deepcopy(STORY), current)
        edited = manifest.validate_manifest(concierge_manifest(), CATALOG, 16)  # the editor dropped the beats
        manifest.attach_models(edited, CATALOG)
        edited.append(dict(edited[-1], id="korean", beat="b9"))
        build_editor._thread_story(edited, {"story": story, "deliverables": current})
        self.assertEqual([d["beat"] for d in edited], ["b2", "b2", "b2", ""])
        self.assertEqual(dlv.spec_hashes(edited[:3]), dlv.spec_hashes(current))


if __name__ == "__main__":
    unittest.main()
