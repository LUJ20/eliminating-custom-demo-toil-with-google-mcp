"""Pre-built sample demos (engine/prebuild.py): when a saved demo stands in for a build, which samples a run
rebuilds, and the triggers (model promotion, app start). Offline: build() is replaced by a fake that writes a
result file the way a real build does."""
import os
import threading
import time
from unittest import mock

import pptx

from engine import model_resolver as mr
from engine import prebuild
from engine.common import read_json, write_json
from engine.deck_generator import DECK_VERSION, deck_version
from engine.usecase_synthesizer import RESULT_FILE
from fakes import OfflineTestCase, seed_registry

CUSTOMER = "Cymbal Air"
ASK = "Multilingual live concierge that answers flight questions by voice and shows the gate on a map."
CURRENT_MODELS = {"reasoning": {"model": "gemini-3.1-pro-preview", "location": "global"},
                  "fast": {"model": "gemini-3.8-flash", "location": "global"}}


def save_result(settings, customer=CUSTOMER, ask=ASK, final_status="PASS", mode=None, models=None) -> str:
    """A finished build's result file in the customer's project folder, as the synthesizer writes it."""
    folder = prebuild.project_dir(settings, customer)
    os.makedirs(folder, exist_ok=True)
    write_json(os.path.join(folder, RESULT_FILE),
               {"customer_name": customer, "usecase_ask": ask, "final_status": final_status,
                "mode": settings.mode if mode is None else mode, "score": 0.9,
                "models": CURRENT_MODELS if models is None else models, "project_dir": folder})
    return folder


class StalenessTest(OfflineTestCase):
    def setUp(self):
        super().setUp()
        seed_registry(self.settings)

    def test_missing_without_a_saved_project(self):
        self.assertEqual("missing", prebuild.staleness(self.settings, CUSTOMER, ASK))
        self.assertIsNone(prebuild.current(self.settings, CUSTOMER, ASK))

    def test_same_ask_on_the_same_models_is_current(self):
        folder = save_result(self.settings)
        self.assertEqual("", prebuild.staleness(self.settings, CUSTOMER, ASK))
        self.assertEqual(folder, prebuild.current(self.settings, CUSTOMER, ASK))

    def test_whitespace_only_differences_do_not_count(self):
        save_result(self.settings)
        spaced = "  " + ASK.replace(" that ", "   that\n") + " \n"
        self.assertEqual("", prebuild.staleness(self.settings, CUSTOMER.upper().lower().title(), spaced))

    def test_any_change_to_the_text_is_a_different_ask(self):
        save_result(self.settings)
        self.assertEqual("different ask", prebuild.staleness(self.settings, CUSTOMER, ASK.replace("map", "chart")))
        self.assertEqual("different ask", prebuild.staleness(self.settings, CUSTOMER, ASK + " Also SMS."))
        self.assertEqual("different ask", prebuild.staleness(self.settings, CUSTOMER, ASK[:-1]))

    def test_unfinished_and_other_mode(self):
        save_result(self.settings, final_status="")
        self.assertEqual("unfinished", prebuild.staleness(self.settings, CUSTOMER, ASK))
        other = "production" if self.settings.mode != "production" else "showcase"
        save_result(self.settings, mode=other)
        self.assertEqual(f"built in {other} mode", prebuild.staleness(self.settings, CUSTOMER, ASK))

    def test_newer_model_in_use_makes_it_stale(self):
        save_result(self.settings, models={"reasoning": {"model": "gemini-3.1-pro-preview", "location": "global"},
                                           "fast": {"model": "gemini-2.5-flash", "location": "global"}})
        self.assertEqual("newer model in use for fast: gemini-2.5-flash -> gemini-3.8-flash",
                         prebuild.staleness(self.settings, CUSTOMER, ASK))
        self.assertIsNone(prebuild.current(self.settings, CUSTOMER, ASK))

    def test_a_tier_without_a_model_right_now_keeps_the_demo(self):
        save_result(self.settings, models={**CURRENT_MODELS, "video": {"model": "veo-3.1", "location": "global"}})
        self.assertEqual("", prebuild.staleness(self.settings, CUSTOMER, ASK))

    def test_status_covers_every_sample(self):
        st = prebuild.status(self.settings)
        self.assertEqual(set(prebuild.SAMPLES), set(st))
        self.assertTrue(all(v == "missing" for v in st.values()))


class RunTest(OfflineTestCase):
    def setUp(self):
        super().setUp()
        seed_registry(self.settings)
        self.cases = [{"name": "Airline", "customer": CUSTOMER, "ask": ASK},
                      {"name": "Retail", "customer": "Cymbal Shops", "ask": "Virtual try-on with a stylist chat."}]
        self.built = []
        self.seen_building = {}

    def fake_build(self, settings, customer, ask, progress=None):
        self.seen_building[customer] = prebuild.building(customer)
        if progress:
            progress("planning")
        self.built.append(customer)
        save_result(settings, customer, ask)
        return {"score": 0.91, "final_status": "PASS", "seconds": 1.5}

    def test_builds_only_the_missing_or_stale_samples(self):
        save_result(self.settings)  # the airline sample is current already
        with mock.patch.object(prebuild, "build", side_effect=self.fake_build):
            report = prebuild.run(self.settings, self.cases, log=lambda m: None)
        self.assertEqual(["Cymbal Shops"], self.built)
        self.assertEqual(["Airline"], report["current"])
        self.assertFalse(report["skipped"])
        row, = report["built"]
        self.assertEqual(("Retail", "missing", 0.91, "PASS", ""),
                         (row["name"], row["reason"], row["score"], row["final_status"], row["error"]))
        self.assertEqual({"reasoning": "gemini-3.1-pro-preview", "fast": "gemini-3.8-flash"}, report["models"])
        self.assertEqual(report["built"], read_json(prebuild.report_path(self.settings), {})["built"])
        self.assertTrue(self.seen_building["Cymbal Shops"], "building() must be true while the sample builds")
        self.assertFalse(prebuild.building("Cymbal Shops"))
        # Both samples now stand in for their asks.
        self.assertEqual([], [c for c in self.cases if prebuild.staleness(self.settings, c["customer"], c["ask"])])

    def test_force_rebuilds_everything(self):
        save_result(self.settings)
        with mock.patch.object(prebuild, "build", side_effect=self.fake_build):
            report = prebuild.run(self.settings, self.cases, force=True, log=lambda m: None)
        self.assertEqual({CUSTOMER, "Cymbal Shops"}, set(self.built))
        self.assertEqual({"forced"}, {r["reason"] for r in report["built"]})

    def test_a_failing_build_is_a_row_not_an_exception(self):
        def boom(settings, customer, ask, progress=None):
            if customer == CUSTOMER:
                raise RuntimeError("Veo quota exhausted for jane.tester@corp.test")
            return self.fake_build(settings, customer, ask, progress)

        with mock.patch.object(prebuild, "build", side_effect=boom):
            report = prebuild.run(self.settings, self.cases, log=lambda m: None)
        rows = {r["name"]: r for r in report["built"]}
        self.assertIn("RuntimeError", rows["Airline"]["error"])
        self.assertNotIn("jane.tester", rows["Airline"]["error"], "errors are redacted in the report")
        self.assertEqual("", rows["Retail"]["error"])
        self.assertEqual("missing", prebuild.staleness(self.settings, CUSTOMER, ASK))

    def test_samples_build_in_parallel_and_wait_joins_a_build(self):
        gate = threading.Event()
        started = []

        def slow(settings, customer, ask, progress=None):
            started.append(customer)
            gate.wait(5)
            return self.fake_build(settings, customer, ask, progress)

        with mock.patch.object(prebuild, "build", side_effect=slow):
            th = threading.Thread(target=prebuild.run, args=(self.settings, self.cases), kwargs={"log": lambda m: None})
            th.start()
            deadline = time.monotonic() + 5
            while len(started) < 2 and time.monotonic() < deadline:
                time.sleep(0.01)
            self.assertEqual(2, len(started), "two samples build at once (MAX_PARALLEL)")
            self.assertTrue(prebuild.building(CUSTOMER))
            self.assertFalse(prebuild.wait(CUSTOMER, timeout=0.05), "wait() times out while the build runs")
            gate.set()
            th.join(5)
        self.assertFalse(th.is_alive())
        self.assertTrue(prebuild.wait(CUSTOMER, timeout=0.05))
        self.assertFalse(prebuild.building(CUSTOMER))

    def test_a_second_run_at_the_same_time_is_skipped(self):
        with prebuild._try_run_lock(self.settings) as acquired:
            self.assertTrue(acquired)
            with mock.patch.object(prebuild, "build", side_effect=self.fake_build):
                report = prebuild.run(self.settings, self.cases, log=lambda m: None)
        self.assertTrue(report["skipped"])
        self.assertEqual([], self.built)

    def test_no_cases_builds_nothing(self):
        with mock.patch.object(prebuild, "build", side_effect=self.fake_build):
            report = prebuild.run(self.settings, [], log=lambda m: None)
        self.assertEqual([], report["built"])
        self.assertEqual([], self.built)


class TriggerTest(OfflineTestCase):
    def setUp(self):
        super().setUp()
        seed_registry(self.settings)
        self.addCleanup(mr.remove_promotion_hook, prebuild.after_promotion)

    def test_promotion_hook_rebuilds_the_samples(self):
        with mock.patch.object(prebuild, "run", return_value={"built": []}) as run:
            prebuild.after_promotion(self.settings, [{"tier": "fast", "model": "gemini-4-flash", "event": "promoted"}],
                                     log=lambda m: None)
        run.assert_called_once()
        self.assertIs(self.settings, run.call_args.args[0])

    def test_promotion_hook_never_raises(self):
        with mock.patch.object(prebuild, "run", side_effect=RuntimeError("disk full")):
            prebuild.after_promotion(self.settings, [{"tier": "fast"}], log=lambda m: None)  # no exception

    def test_promotion_hook_respects_the_switch(self):
        with mock.patch.object(prebuild, "enabled", return_value=False), \
                mock.patch.object(prebuild, "run") as run:
            prebuild.after_promotion(self.settings, [{"tier": "fast"}], log=lambda m: None)
        run.assert_not_called()

    def test_install_registers_the_hook_once(self):
        self.assertTrue(prebuild.install(self.settings))
        self.assertTrue(prebuild.install(self.settings))
        self.assertEqual(1, mr._PROMOTION_HOOKS.count(prebuild.after_promotion))

    def test_start_background_runs_once_models_are_resolved(self):
        done = threading.Event()

        def fake_run(settings, *a, **k):
            done.set()
            return {"built": []}

        with mock.patch.object(prebuild, "run", side_effect=fake_run), \
                mock.patch.object(prebuild, "POLL_S", 0):
            th = prebuild.start_background(self.settings)
            self.assertIsNotNone(th)
            self.assertTrue(done.wait(5))
            th.join(5)

    def test_start_background_is_off_with_the_switch(self):
        with mock.patch.object(prebuild, "enabled", return_value=False):
            self.assertIsNone(prebuild.start_background(self.settings))


class CliTest(OfflineTestCase):
    def setUp(self):
        super().setUp()
        seed_registry(self.settings)

    def test_status_and_sample_filter(self):
        with mock.patch.object(prebuild, "get_settings", return_value=self.settings), \
                mock.patch("builtins.print") as out:
            self.assertEqual(0, prebuild.main(["--status"]))
        printed = "\n".join(str(c.args[0]) for c in out.call_args_list)
        for name in prebuild.SAMPLES:
            self.assertIn(name, printed)
        with mock.patch.object(prebuild, "get_settings", return_value=self.settings):
            with self.assertRaises(SystemExit):
                prebuild.main(["--sample", "No Such Customer"])

    def test_main_builds_and_reports(self):
        first = next(iter(prebuild.SAMPLES.values()))

        def fake_build(settings, customer, ask, progress=None):
            save_result(settings, customer, ask)
            return {"score": 0.95, "final_status": "PASS", "seconds": 2}

        with mock.patch.object(prebuild, "get_settings", return_value=self.settings), \
                mock.patch.object(prebuild, "build", side_effect=fake_build), \
                mock.patch("builtins.print"):
            self.assertEqual(0, prebuild.main(["--sample", first["customer"]]))
        self.assertEqual("", prebuild.staleness(self.settings, first["customer"], first["ask"]))


class DeckRefreshTest(OfflineTestCase):
    """A new slide layout reaches the saved demos without a rebuild (build_editor.refresh_deck)."""

    def setUp(self):
        super().setUp()
        seed_registry(self.settings)

    def finished_project(self, customer=CUSTOMER, stages=True) -> str:
        folder = save_result(self.settings, customer)
        res = read_json(os.path.join(folder, RESULT_FILE), {})
        res.update(summary="One concierge.", package_files=["pipeline.py"], whats_new=[], deliverables=[],
                   eval_metrics=[{"metric": "Code validity", "value": "PASS", "threshold": "compiles",
                                  "method": "programmatic", "notes": "", "pass": True}],
                   attempt_stats=[], stages=[{"stage": "1. Ingest", "service": "Cloud Run", "api": "run", "tier": "",
                                              "model": "", "description": "d", "features": []}] if stages else [])
        write_json(os.path.join(folder, RESULT_FILE), res)
        write_json(os.path.join(folder, "usecase_config.json"), {"summary": "One concierge.", "stages": res["stages"]})
        with open(os.path.join(folder, "pipeline.py"), "w", encoding="utf-8") as f:
            f.write("print('hi')\n")
        pptx.Presentation().save(os.path.join(folder, f"{os.path.basename(folder)}_architecture_deck.pptx"))
        return folder

    def test_an_older_deck_is_regenerated_from_the_stored_result(self):
        folder = self.finished_project()
        deck = os.path.join(folder, f"{os.path.basename(folder)}_architecture_deck.pptx")
        self.assertEqual("", deck_version(deck))
        self.assertEqual([os.path.basename(folder)], prebuild.refresh_decks(self.settings, log=lambda m: None))
        self.assertEqual(DECK_VERSION, deck_version(deck))
        prs = pptx.Presentation(deck)
        self.assertEqual(5, len(prs.slides))
        text = " ".join(sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame)
        self.assertIn("Cloud Run", text)
        self.assertEqual([], prebuild.refresh_decks(self.settings, log=lambda m: None))  # already current

    def test_a_result_without_a_design_is_left_alone(self):
        folder = self.finished_project(stages=False)
        deck = os.path.join(folder, f"{os.path.basename(folder)}_architecture_deck.pptx")
        self.assertEqual([], prebuild.refresh_decks(self.settings, log=lambda m: None))
        self.assertEqual("", deck_version(deck))

    def test_decks_only_cli_does_not_build(self):
        self.finished_project()
        with mock.patch.object(prebuild, "get_settings", return_value=self.settings), \
                mock.patch.object(prebuild, "build") as build, mock.patch("builtins.print") as out:
            self.assertEqual(0, prebuild.main(["--decks"]))
        build.assert_not_called()
        self.assertIn("1 deck(s) regenerated", "\n".join(str(c.args[0]) for c in out.call_args_list))

    def test_start_background_refreshes_decks_before_building(self):
        order = []
        with mock.patch.object(prebuild, "refresh_decks", side_effect=lambda *a, **k: order.append("decks") or []), \
                mock.patch.object(prebuild, "run", side_effect=lambda *a, **k: order.append("run") or {"built": []}), \
                mock.patch.object(prebuild, "POLL_S", 0):
            th = prebuild.start_background(self.settings)
            th.join(5)
        self.assertEqual(["decks", "run"], order)


class ChatRefreshTest(OfflineTestCase):
    """A chat set up before the assistant got its context is directed and played again in place
    (deliverables.redirect through prebuild.refresh_chats), with no rebuild of the project."""
    finished_project = DeckRefreshTest.finished_project

    def setUp(self):
        super().setUp()
        seed_registry(self.settings)

    def _status(self, folder: str, with_reply: bool) -> dict:
        from engine import deliverables as dlv
        os.makedirs(dlv.folder(folder), exist_ok=True)
        asset = {"label": "English", "language": "en", "status": "ready", "prompt": "You are the concierge.",
                 "script": "Hi!", "file": "", "mime": "", "model": "m", "seconds": 1, "error": "", "note": "",
                 "qa": {"verdict": "pass"}, "spec_hash": "h", "carried_over": False}
        if with_reply:
            asset.update(file="chat__english__abcd.md", mime="text/markdown")
            with open(os.path.join(dlv.folder(folder), asset["file"]), "w", encoding="utf-8") as f:
                f.write("Welcome back, Aiko.")
        state = {"build_id": "b1", "job_id": "j1", "state": "done", "error": "", "customer": CUSTOMER, "ask": "a",
                 "summary": "s", "story": {},
                 "deliverables": [{"id": "chat", "title": "Try the concierge", "kind": "chat", "tier": "fast",
                                   "brief": "Ask.", "start_from": "", "model": "", "location": "", "status": "",
                                   "beat": "", "scene": "", "spec_hash": "h", "assets": [asset]}]}
        write_json(dlv.status_path(folder), state)
        return state

    def test_only_chats_without_a_played_reply_are_redirected(self):
        from engine import brain, deliverables as dlv, media_qa
        old = self.finished_project()
        self._status(old, with_reply=False)
        new = self.finished_project("Other Co")
        self._status(new, with_reply=True)
        directed, turns = [], []

        def direction(*args, deliverables, **kwargs):
            directed.extend((d["id"], v["label"], d["kind"]) for d in deliverables for v in d["variants"])
            return {(d["id"], v["label"]): {"prompt": "You are the concierge.\nContext: Aiko, flight DL 275.",
                                            "script": "Where is my flight?"}
                    for d in deliverables for v in d["variants"]}, 1.0

        def chat(settings, model, location, system, history):
            turns.append((system, [h["text"] for h in history]))
            return "DL 275 leaves at 09:40, Aiko."

        with mock.patch.object(brain, "direct_media", side_effect=direction), \
                mock.patch.object(dlv.media, "chat", side_effect=chat), \
                mock.patch.object(media_qa, "check", side_effect=lambda *a, **k: {
                    "verdict": "pass", "score": 1.0, "summary": "fine", "checks": []}), \
                mock.patch.object(dlv, "McpKnowledgeClient", return_value=mock.MagicMock()):
            done = prebuild.refresh_chats(self.settings, log=lambda m: None)
            for t in list(dlv._JOBS.values()):
                t.join(10)
        self.assertEqual([os.path.basename(old)], done)
        self.assertEqual([("chat", "English", "chat")], directed)  # the setup is written again, once
        self.assertEqual([("You are the concierge.\nContext: Aiko, flight DL 275.", ["Where is my flight?"])], turns)
        asset = dlv.load_status(old)["deliverables"][0]["assets"][0]
        self.assertEqual(("ready", "Where is my flight?"), (asset["status"], asset["script"]))
        self.assertEqual("DL 275 leaves at 09:40, Aiko.", dlv.chat_reply(old, asset))
        untouched = dlv.load_status(new)["deliverables"][0]["assets"][0]
        self.assertEqual(("Hi!", "Welcome back, Aiko."), (untouched["script"], dlv.chat_reply(new, untouched)))
        self.assertEqual([], prebuild.refresh_chats(self.settings, log=lambda m: None))  # nothing left to redo
