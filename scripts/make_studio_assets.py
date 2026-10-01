"""Regenerate the Studio's own docs assets from the current architecture and live project state:

  docs/studio_architecture_flow.svg   architecture diagram (vector)
  docs/studio_architecture_flow.png   the same diagram as PNG (2x)
  docs/Studio_Architecture_Deck.pptx  editable studio deck (native shapes and tables)

The use-case slide reads the saved projects in the output folder and the model slide reads the model registry, so
re-running this script after builds or model upgrades keeps the deck current. Usage:

    .venv/bin/python scripts/make_studio_assets.py
"""
import datetime as dt
import json
import os
import sys
from html import escape

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
from engine.config import get_settings  # noqa: E402
from engine.model_resolver import ModelResolver  # noqa: E402

DOCS = os.path.join(ROOT, "docs")
SVG_PATH = os.path.join(DOCS, "studio_architecture_flow.svg")
PNG_PATH = os.path.join(DOCS, "studio_architecture_flow.png")
PPTX_PATH = os.path.join(DOCS, "Studio_Architecture_Deck.pptx")
TITLE = "Gemini + MCP Use-Case Studio"
RESULT_FILE = ".studio_result.json"

BLUE, GREEN, YELLOW, RED, PURPLE = "#1A73E8", "#188038", "#F9AB00", "#D93025", "#A142F4"
INK, GREY, LINE, PAPER = "#202124", "#5F6368", "#DADCE0", "#FFFFFF"
# Fonts with an arrow glyph (Helvetica on macOS has none); the first one found is used.
FONT_FILES = ("/System/Library/Fonts/Supplemental/Arial.ttf", "/Library/Fonts/Arial.ttf",
              "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
BOLD_FILES = ("/System/Library/Fonts/Supplemental/Arial Bold.ttf", "/Library/Fonts/Arial Bold.ttf",
              "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")

# ------------------------------------------------------------------ diagram content (one place to edit)
FLOW = ["Ask", "Ground (MCP)", "Plan + story", "Code", "Judge + fix", "Acceptance tests", "Outputs + checks",
        "Deliver"]
LANES = [
    ("EXPERIENCE", "Streamlit · app.py", BLUE, [
        ("Ask", "Customer and use case in plain words; sample presets across use-case types"),
        ("Saved projects", "Open any stored build with its scorecard, outputs, chat and versions"),
        ("Result view", "Architecture · demo outputs · code + PII audit · scorecard · story, deck, codebase"),
        ("Build chat", "Ask anything (answers cite official docs) or change the build: undo, save, discard"),
        ("Settings", "Showcase or Production models · Drive folder · model events with reasons"),
    ]),
    ("ORCHESTRATOR", "Python workflow · decides the steps", GREEN, [
        ("Ground", "Search official Google docs for the ask through MCP"),
        ("Plan", "Stages on any Google service, demo outputs and a story (reasoning tier)"),
        ("Code", "pipeline.py with --dry-run; model IDs read from config (fast tier)"),
        ("Judge + fix", "Scorecard; failed rows fed back as fixes; up to 3 attempts"),
        ("Acceptance tests", "3-6 tests written from the ask, run on the chosen models"),
        ("Outputs job", "Starts at the first valid plan; 6 in parallel; check, regenerate, keep best"),
        ("Chat editor", "Question or change; code edits re-validated, re-judged, re-scored"),
    ]),
    ("BRAIN · KNOWLEDGE · TOOLS", "Gemini · MCP · Google APIs", YELLOW, [
        ("Gemini brain", "Plans, writes code, judges, checks outputs, answers chat; reasoning, fast and lite tiers; validated JSON"),
        ("Developer Knowledge MCP", "The knowledge source: search_documents, get_documents; every stage, model and answer cites a doc"),
        ("Google API tools", "Called directly by code: Veo, Gemini image, Gemini-TTS, Lyria, Live, embeddings, Drive, Cloud Storage"),
        ("Model Resolver", "Picks the brain's newest verified model per tier from the docs; never hardcoded"),
        ("Troubleshooter", "Classifies failures, applies safe fixes, retries, falls back; MCP-grounded diagnosis"),
    ]),
    ("EVALS", "6 layers", RED, [
        ("1  Build scorecard", "5 judge rows (each 4/5 or more) + 6-7 programmatic checks"),
        ("2  Acceptance tests", "Use case works end to end: 80% pass, no safety failure"),
        ("3  Output checks", "Per output type: language, lip-sync, schema, safe actions, brand safety"),
        ("4  Chat checks", "Citations verified against docs; change done / partly / not done"),
        ("5  Model gates", "9-task golden set · media canary · daily drift · runtime watch"),
        ("6  Regression", "4 reference use cases rebuilt after any model upgrade"),
    ]),
    ("DELIVERABLES", "per build", PURPLE, [
        ("Reference architecture", "Stages, services, models and doc citations"),
        ("Demo outputs", "Video with lip-sync, image, speech, music, text, chat, JSON result, agent trace"),
        ("Code", "pipeline.py and a PII-scanned codebase zip"),
        ("Scorecard", "Every row with the judge's reasons; PASSED or BEST EFFORT"),
        ("Story script + deck", "Google Doc and Slides in Drive mode, else .html and .pptx"),
        ("Publish", "Drive folder or Cloud Storage"),
    ]),
]
BANDS = [
    (0, 2, "State", "generated_projects/<slug>: code, outputs, versions  ·  .cache: model registry, telemetry, "
                    "regression baselines"),
    (2, 5, "Self-upgrading models", "discover → verify → golden set or canary → promote → watch → regression → "
                                    "roll back or forward (24 h quarantine); failed output checks count against "
                                    "the model"),
]


# ------------------------------------------------------------------ drawing primitives (SVG + PNG from one list)
def _font(size: float, bold: bool = False):
    for path in (BOLD_FILES if bold else FONT_FILES):
        if os.path.exists(path):
            try:
                return ImageFont.truetype(path, int(round(size)), index=1 if bold and path.endswith(".ttc") else 0)
            except OSError:
                continue
    return ImageFont.load_default()


def _text_width(text: str, size: float, bold: bool = False) -> float:
    return _font(size * 4, bold).getlength(text) / 4  # measured at 4x: font sizes are whole pixels


def _wrap(text: str, size: float, bold: bool, width: float) -> list:
    lines, cur = [], ""
    for word in text.split():
        trial = f"{cur} {word}".strip()
        if _text_width(trial, size, bold) <= width or not cur:
            cur = trial
        else:
            lines.append(cur)
            cur = word
    return lines + ([cur] if cur else [])


def _rgb(hex_color: str) -> tuple:
    return tuple(int(hex_color[i:i + 2], 16) for i in (1, 3, 5))


def _tint(hex_color: str, amount: float) -> str:
    r, g, b = _rgb(hex_color)
    return "#%02X%02X%02X" % tuple(int(c + (255 - c) * amount) for c in (r, g, b))


def build_diagram() -> tuple:
    """-> (width, height, primitives)."""
    W, lane_w, gap, x0 = 1600, 278, 37, 34
    ops = [("rect", 0, 0, W, 0, PAPER, PAPER, 0)]  # height patched below
    ops.append(("text", x0, 46, f"{TITLE}: architecture", 30, True, INK, "start"))
    ops.append(("text", x0, 74, "A Python workflow orchestrates a Gemini brain, the Developer Knowledge MCP server "
                                "(official docs) and Google APIs as tools. Every step is evaluated.", 15, False, GREY,
                "start"))
    # flow strip
    fx, fy, fh = x0, 96, 34
    fw = (W - 2 * x0 - (len(FLOW) - 1) * 6) / len(FLOW)
    for i, step in enumerate(FLOW):
        x = fx + i * (fw + 6)
        pts = [(x, fy), (x + fw - 12, fy), (x + fw, fy + fh / 2), (x + fw - 12, fy + fh), (x, fy + fh)]
        if i:
            pts.append((x + 12, fy + fh / 2))
        ops.append(("poly", pts, _tint(GREY, 0.85) if i % 2 else _tint(GREY, 0.78)))
        ops.append(("text", x + fw / 2 + (6 if i else 0), fy + 22, step, 13, True, INK, "middle"))
    # lanes
    top, head_h, pad, box_gap = 150, 50, 12, 10
    lane_boxes, lane_bottom = [], top
    for li, (name, sub, color, boxes) in enumerate(LANES):
        lx, y, laid = x0 + li * (lane_w + gap), top + head_h + pad, []
        for title, body in boxes:
            tl = _wrap(title, 14, True, lane_w - 2 * pad - 24)
            bl = _wrap(body, 12.5, False, lane_w - 2 * pad - 24)
            h = 14 + 18 * len(tl) + 16 * len(bl) + 6
            laid.append((lx + pad, y, lane_w - 2 * pad, h, tl, bl))
            y += h + box_gap
        lane_boxes.append((lx, color, name, sub, laid))
        lane_bottom = max(lane_bottom, y + pad - box_gap)
    for lx, color, name, sub, laid in lane_boxes:
        ops.append(("rect", lx, top, lane_w, lane_bottom - top, _tint(color, 0.93), _tint(color, 0.6), 10))
        ops.append(("rect", lx, top, lane_w, head_h, color, color, 10))
        head_ink = INK if color == YELLOW else PAPER
        ops.append(("text", lx + lane_w / 2, top + 22, name, 15, True, head_ink, "middle"))
        ops.append(("text", lx + lane_w / 2, top + 40, sub, 12, False, head_ink, "middle"))
        for bx, by, bw, bh, tl, bl in laid:
            ops.append(("rect", bx, by, bw, bh, PAPER, _tint(color, 0.45), 7))
            ty = by + 20
            for line in tl:
                ops.append(("text", bx + 10, ty, line, 14, True, INK, "start"))
                ty += 18
            for line in bl:
                ops.append(("text", bx + 10, ty, line, 12.5, False, GREY, "start"))
                ty += 16
    # arrows between lanes
    for li in range(len(LANES) - 1):
        ax = x0 + (li + 1) * lane_w + li * gap
        for ay in (top + head_h + 60, (top + lane_bottom) / 2 + 40):
            ops.append(("line", ax + 5, ay, ax + gap - 9, ay, GREY, 2.5))
            ops.append(("poly", [(ax + gap - 4, ay), (ax + gap - 13, ay - 6), (ax + gap - 13, ay + 6)], GREY))
    # bottom bands
    by = lane_bottom + 18
    for a, b, title, text in BANDS:
        bx = x0 + a * (lane_w + gap)
        bw = (b - a) * lane_w + (b - a - 1) * gap
        lines = _wrap(text, 13, False, bw - 40 - _text_width(title, 14, True))
        bh = 20 + 17 * len(lines)
        ops.append(("rect", bx, by, bw, bh, _tint(GREY, 0.92), _tint(GREY, 0.55), 8))
        ops.append(("text", bx + 12, by + 23, title, 14, True, INK, "start"))
        tx = bx + 24 + _text_width(title, 14, True)
        for i, line in enumerate(lines):
            ops.append(("text", tx, by + 23 + 17 * i, line, 13, False, INK, "start"))
    H = int(by + 20 + 17 * 3 + 30)
    ops[0] = ("rect", 0, 0, W, H, PAPER, PAPER, 0)
    ops.append(("text", W - x0, H - 12, f"generated by scripts/make_studio_assets.py · {dt.date.today()}", 11, False,
                GREY, "end"))
    return W, H, ops


def write_svg(W: int, H: int, ops: list, path: str) -> None:
    out = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}" viewBox="0 0 {W} {H}" '
           'font-family="Arial, Helvetica, sans-serif">']
    for op in ops:
        kind = op[0]
        if kind == "rect":
            _, x, y, w, h, fill, stroke, r = op
            out.append(f'<rect x="{x:.1f}" y="{y:.1f}" width="{w:.1f}" height="{h:.1f}" rx="{r}" fill="{fill}" '
                       f'stroke="{stroke}" stroke-width="1.2"/>')
        elif kind == "text":
            _, x, y, text, size, bold, color, anchor = op
            weight = ' font-weight="bold"' if bold else ""
            out.append(f'<text x="{x:.1f}" y="{y:.1f}" font-size="{size}" fill="{color}" text-anchor="{anchor}"'
                       f'{weight}>{escape(text)}</text>')
        elif kind == "line":
            _, x1, y1, x2, y2, color, width = op
            out.append(f'<line x1="{x1:.1f}" y1="{y1:.1f}" x2="{x2:.1f}" y2="{y2:.1f}" stroke="{color}" '
                       f'stroke-width="{width}"/>')
        elif kind == "poly":
            _, pts, fill = op
            out.append(f'<polygon points="{" ".join(f"{x:.1f},{y:.1f}" for x, y in pts)}" fill="{fill}"/>')
    out.append("</svg>")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(out) + "\n")


def write_png(W: int, H: int, ops: list, path: str, scale: int = 2) -> None:
    img = Image.new("RGB", (W * scale, H * scale), PAPER)
    d = ImageDraw.Draw(img)
    anchors = {"start": "ls", "middle": "ms", "end": "rs"}
    for op in ops:
        kind = op[0]
        if kind == "rect":
            _, x, y, w, h, fill, stroke, r = op
            d.rounded_rectangle([x * scale, y * scale, (x + w) * scale, (y + h) * scale], radius=r * scale,
                                fill=fill, outline=stroke, width=max(1, int(1.2 * scale)))
        elif kind == "text":
            _, x, y, text, size, bold, color, anchor = op
            d.text((x * scale, y * scale), text, font=_font(size * scale, bold), fill=color, anchor=anchors[anchor])
        elif kind == "line":
            _, x1, y1, x2, y2, color, width = op
            d.line([x1 * scale, y1 * scale, x2 * scale, y2 * scale], fill=color, width=int(width * scale))
        elif kind == "poly":
            _, pts, fill = op
            d.polygon([(x * scale, y * scale) for x, y in pts], fill=fill)
    img.save(path, optimize=True)


# ------------------------------------------------------------------ live facts for the deck
def saved_projects(output_dir: str) -> list:
    rows = []
    for name in sorted(os.listdir(output_dir)) if os.path.isdir(output_dir) else []:
        path = os.path.join(output_dir, name, RESULT_FILE)
        try:
            with open(path, encoding="utf-8") as f:
                r = json.load(f)
        except (OSError, ValueError):
            continue
        acc = r.get("acceptance") or {}
        rows.append({
            "customer": r.get("customer_name") or name,
            "ask": r.get("usecase_ask", ""),
            "kinds": ", ".join(sorted({d.get("kind", "") for d in r.get("deliverables", []) if d.get("kind")})),
            "score": r.get("score"),
            "status": r.get("final_status", ""),
            "e2e": f"{acc['passed']}/{acc['total']}" if acc.get("total") else "not run",
            "at": os.path.getmtime(path),
        })
    return sorted(rows, key=lambda x: -x["at"])


def champions(settings) -> list:
    path = ModelResolver(settings).path
    try:
        with open(path, encoding="utf-8") as f:
            tiers = json.load(f).get("tiers", {})
    except (OSError, ValueError):
        return []
    return [(t, (v.get("champion") or {}).get("model", "")) for t, v in tiers.items() if v.get("champion")]


def _short(text: str, n: int) -> str:
    text = " ".join(str(text).split())
    return text if len(text) <= n else text[:n].rsplit(" ", 1)[0] + "…"


# ------------------------------------------------------------------ deck
def _color(hex_color: str) -> RGBColor:
    return RGBColor(*_rgb(hex_color))


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.core_properties.author = self.prs.core_properties.last_modified_by = "Gemini + MCP Use-Case Studio"
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.n = 0

    def slide(self, title: str, subtitle: str = ""):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.n += 1
        self.text(s, title, 0.5, 0.35, 12.3, 0.7, 28, True, INK)
        if subtitle:
            self.text(s, subtitle, 0.5, 0.95, 12.3, 0.45, 15, False, GREY)
        self.text(s, f"{TITLE}  |  {self.n}", 0.5, 7.0, 12.3, 0.3, 10, False, GREY)
        return s

    @staticmethod
    def text(s, text, x, y, w, h, size, bold=False, color=INK, align=PP_ALIGN.LEFT):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = line
            r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, _color(color)
        return tb

    @staticmethod
    def box(s, x, y, w, h, fill, title, body="", title_color=PAPER, body_color=PAPER, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
            title_size=15):
        sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.fill.solid()
        sh.fill.fore_color.rgb = _color(fill)
        sh.line.fill.background()
        tf = sh.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = title
        r.font.size, r.font.bold, r.font.color.rgb = Pt(title_size), True, _color(title_color)
        if body:
            p2 = tf.add_paragraph()
            r2 = p2.add_run()
            r2.text = body
            r2.font.size, r2.font.color.rgb = Pt(11.5), _color(body_color)
        return sh

    @staticmethod
    def table(s, rows, x, y, w, col_w, size=12):
        shape = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(0.4 * len(rows)))
        t = shape.table
        for ci, cw in enumerate(col_w):
            t.columns[ci].width = Emu(int(Inches(cw)))
        for ri, row in enumerate(rows):
            for ci, val in enumerate(row):
                cell = t.cell(ri, ci)
                cell.text = str(val)
                for p in cell.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size, r.font.bold = Pt(size), ri == 0
                        r.font.color.rgb = _color(PAPER if ri == 0 else INK)
                cell.fill.solid()
                cell.fill.fore_color.rgb = _color(BLUE if ri == 0 else ("#F1F3F4" if ri % 2 else PAPER))
        return t

    def save(self, path):
        self.prs.save(path)


def build_deck(settings, png_path: str, path: str) -> None:
    d = Deck()
    projects = saved_projects(settings.output_dir)
    champs = champions(settings)

    s = d.slide(TITLE, "Any customer use case → a grounded, current, story-driven, evaluated demo in minutes")
    stats = [("≈ 7-8 min", "per custom demo package, all outputs checked", BLUE),
             ("0", "manual model updates: the resolver upgrades, gated by evals", GREEN),
             ("6", "eval layers: build, use case, outputs, chat, model gates, regression", RED),
             (str(len(projects)), "saved projects across use-case types", PURPLE)]
    for i, (big, small, color) in enumerate(stats):
        d.box(s, 0.5 + i * 3.1, 2.0, 2.9, 2.2, color, big, small, title_size=36)
    d.text(s, "Grounded in official Google docs through the Developer Knowledge MCP server · newest verified models, "
              "never hardcoded · real outputs that tell a story · code, scorecard, story script and editable deck",
           0.5, 4.7, 12.3, 1.0, 15, False, GREY)

    s = d.slide("Architecture", "Experience → build → grounding and models → evals → deliverables")
    with Image.open(png_path) as im:
        ratio = im.width / im.height
    h = 5.75
    w = min(12.3, h * ratio)
    s.shapes.add_picture(png_path, Inches((13.333 - w) / 2), Inches(1.15), Inches(w), Inches(w / ratio))

    s = d.slide("Brain, knowledge and tools", 'Not "Gemini → MCP → all tools": the orchestrator calls each part')
    d.box(s, 0.5, 1.5, 12.3, 0.9, GREY, "Orchestrator: Python workflow",
          "Decides the order of steps, retries, budgets and when to stop (build, chat editor, outputs job)")
    parts = [("Brain: Gemini", "Plans, writes code, judges, checks outputs, answers chat. One model per tier, picked "
                               "by the Model Resolver. Always validated JSON.", BLUE),
             ("Knowledge: MCP", "One server, Google Developer Knowledge (search_documents, get_documents). Every "
                                "design, model ID, feature and answer cites an official doc.", YELLOW),
             ("Tools: Google APIs", "Called directly by code: Veo, Gemini image, Gemini-TTS, Lyria, Live, embeddings, "
                                    "Model Garden, Drive, Slides, Docs, Cloud Storage.", GREEN)]
    for i, (t, b, color) in enumerate(parts):
        d.box(s, 0.5 + i * 4.15, 2.7, 4.0, 2.3, color, t, b,
              title_color=INK if color == YELLOW else PAPER, body_color=INK if color == YELLOW else PAPER)
    d.text(s, "Why not ADK: the Studio is a fixed, evaluated workflow, so it needs no agent that picks tools itself. "
              "ADK is worth adding only if chat should act through tools on its own or the Studio moves to a managed "
              "agent runtime. Generated designs can still recommend ADK for a customer's agent use case.",
           0.5, 5.3, 12.3, 1.2, 14, False, INK)

    s = d.slide("How a build flows", "Outputs start at the first valid plan, so clips are made while code is written and judged")
    steps = [("Ask", "customer + use case"), ("Ground", "official docs (MCP)"), ("Plan", "stages, outputs, story"),
             ("Code", "pipeline.py, models from config"), ("Judge", "scorecard, up to 3 attempts"),
             ("Test", "acceptance tests from the ask"), ("Generate", "outputs checked, best kept"),
             ("Deliver", "deck, story script, zip")]
    for i, (t, b) in enumerate(steps):
        d.box(s, 0.4 + i * 1.57, 2.2, 1.62, 1.9, GREEN if i % 2 else BLUE, t, b, shape=MSO_SHAPE.CHEVRON,
              title_size=13)
    d.text(s, "Then: build chat answers questions with verified doc citations, or changes the build as a new version "
              "(undo, save, discard). Code edits are re-validated, re-judged and re-scored with a diff.",
           0.5, 4.7, 12.3, 1.0, 15, False, INK)

    s = d.slide("Any use case, not just media", "Saved projects (newest first), read from the output folder")
    rows = [["Customer", "Use case", "Outputs", "Score", "End to end"]]
    for p in projects[:8]:
        rows.append([p["customer"], _short(p["ask"], 70), p["kinds"],
                     f"{p['score']:.1f}" if isinstance(p["score"], (int, float)) else "-", p["e2e"]])
    d.table(s, rows, 0.5, 1.5, 12.3, [2.6, 4.9, 2.5, 1.0, 1.3], size=11)

    s = d.slide("Six eval layers", "Every build, output, chat change and model upgrade is evaluated (docs/EVAL_RULES.md)")
    d.table(s, [["Layer", "Question", "Pass rule"],
                ["1 Build scorecard", "Are the design and code right?", "Judge rows 4/5 or more; programmatic checks; PASSED = every row"],
                ["2 Acceptance tests", "Does the use case work end to end?", "3-6 tests from the ask; 80% pass, no safety failure"],
                ["3 Output checks", "Is each output right?", "Per type; critical failure → regenerate, best kept"],
                ["4 Chat checks", "Is the answer supported, the change done?", "Citations verified; done / partly / not done"],
                ["5 Model gates", "Is a new model as good, and does it stay good?", "Golden 7/9 or more; media canary; drift ≤ 1 task; watch"],
                ["6 Regression", "Did an upgrade make demos worse?", "4 reference builds; drop > 10 points → roll back"]],
            0.5, 1.5, 12.3, [2.4, 4.4, 5.5], size=12)

    s = d.slide("Models upgrade themselves", "Discovered in docs, verified, tested, promoted, watched, rolled back if worse")
    d.text(s, "1  Discover new model IDs in official docs (MCP)\n2  Verify in Model Garden and with a live call\n"
              "3  Test: 9-task golden set (text) or a real sample (media)\n4  Promote only if at least as good\n"
              "5  Watch every call; failed output checks count\n6  Rebuild 4 reference use cases after an upgrade\n"
              "7  Roll back + 24 h quarantine, or roll forward\n\nIgnored, never a rollback: expired credentials, "
              "network, project setup, rate limits (HTTP 429)", 0.5, 1.5, 6.0, 5.0, 15, False, INK)
    if champs:
        d.table(s, [["Tier", "Current model"]] + [[t, m] for t, m in champs], 7.0, 1.5, 5.8, [2.0, 3.8], size=11)

    s = d.slide("Build chat", "Ask anything, or change the build in plain words")
    d.table(s, [["Ask", "Result"],
                ["A question about the build, Google products, pricing, scaling", "Answer from build facts + MCP docs, verified citations, no version saved"],
                ["Change outputs (\"add Korean\", \"older avatar\")", "New manifest; only changed outputs regenerated"],
                ["Change docs (\"add a cost note\")", "README and deck rewritten"],
                ["Change code (\"add retries to pipeline.py\")", "Rewritten, validated, re-judged, re-scored, diff shown"],
                ["Change the Studio, unverified models, skip security, over the cost cap", "Refused with a reason"]],
            0.5, 1.5, 12.3, [5.4, 6.9], size=12)

    s = d.slide("Status and next steps")
    d.box(s, 0.5, 1.3, 6.0, 4.9, GREEN, "Built",
          "Grounded planning, codegen, judge loop · acceptance tests for any use case · output checks for 8 output "
          "types · storytelling, story script, deck · build chat with citations, change checks, code edits, versions · "
          "self-upgrading models with golden set, canary, watch, drift, regression · saved projects · offline unit "
          "tests and a chat intent set")
    d.box(s, 6.8, 1.3, 6.0, 4.9, BLUE, "Next",
          "Run generated code in an isolated Cloud Run job as part of acceptance · Live API conversation canary · "
          "Cloud Run + IAP, Firestore state · optional: ADK only if chat should act through tools on its own or "
          "the Studio moves to a managed agent runtime")
    d.save(path)


def main() -> None:
    settings = get_settings()
    W, H, ops = build_diagram()
    write_svg(W, H, ops, SVG_PATH)
    write_png(W, H, ops, PNG_PATH)
    build_deck(settings, PNG_PATH, PPTX_PATH)
    for p in (SVG_PATH, PNG_PATH, PPTX_PATH):
        print(f"wrote {os.path.relpath(p, ROOT)} ({os.path.getsize(p) // 1024} KB)")


if __name__ == "__main__":
    main()
